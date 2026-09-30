"""Config-driven Word (.docx) and PDF builder (plan Task 8.2; INV-21..38).

Usage: python harness/tools/build_book.py --project projects/<slug> [--no-pdf]
Reads the chapter files named by chapter-plan.json, brief.json (title, credits, audience line), theme.json (fonts,
palette, callouts, labels, page) and the preset's title-page layout. `run build` writes build/<basename>.docx and,
through the Word COM backend, build/<basename>.pdf; the CLI is a dry run into a temporary folder (step9 fix S9-03:
`run build` is the only writer of build/). Figures: with a figure manifest (<paths.figures>/figures.json, core §7) a
chapter places `![](fig:<id>)`; the raster comes from <paths.figures>/out/ (rendered by `run build` first), the
caption and alt text from the manifest and the number is computed. A project without a manifest is a legacy
(history-imported) book: SVG links are read from their PNG mirror in <paths.figures>/png/ and captions from the
chapter. Exit 0 ok; 1 on a refused gate or a build error; 2 on a config error.

Text direction and language go through one factory (make_paragraph / make_run / make_table, EXT-LOC-2). LTR adds no
run-level properties (the language sits on the Normal style); RTL rendering is STEP 10.
"""
import json, math, pathlib, re, sys, tempfile

REPO = pathlib.Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))
from harness import figures  # noqa: E402
from harness.figures import annotated  # noqa: E402
from harness.tools import assemble, blocks, config, mathml  # noqa: E402

# Preset constants of the ltr-textbook layout (sizes in pt, rules and grid colours): tested layout, not project config
RULE, GRID, STRIPE, LINK, WHITE = "F0B429", "C9D1D9", "F4F6F8", "2B5D8A", "FFFFFF"
DEFAULT_BOX = "F1F1EE"
GAP = "\u00a0\u00a0 "   # two no-break spaces and a space between a heading number and its title
LO_ITEM = re.compile(r"\[(LO\d+)\]\s*(.+)")
DOI = re.compile(r"(10\.\d{4,9}/[^\s]+[^\s.,;])")


class BuildError(Exception):
    """A build step failed; the message names it."""


def _docx():
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    return dict(Document=Document, WD_SECTION=WD_SECTION, WD_STYLE_TYPE=WD_STYLE_TYPE, WD_TABLE_ALIGNMENT=WD_TABLE_ALIGNMENT,
                WD_ALIGN_PARAGRAPH=WD_ALIGN_PARAGRAPH, WD_BREAK=WD_BREAK, WD_LINE_SPACING=WD_LINE_SPACING,
                OxmlElement=OxmlElement, qn=qn, Cm=Cm, Pt=Pt, RGBColor=RGBColor)


class Book:
    """Everything the renderer needs, read once from config. No display string is compared in code (EXT-LOC-3)."""

    def __init__(self, cfg):
        self.cfg, t, th, b = cfg, cfg["template"], cfg["theme"], cfg["brief"]
        self.fonts, self.pal = th["fonts"], th["palette"]
        self.serif, self.serif_head, self.sans = self.fonts["serif"], self.fonts["serif_heading"], self.fonts["sans"]
        self.ink, self.primary, self.accent, self.muted = (self.pal[k].lstrip("#") for k in ("ink", "primary", "accent", "muted"))
        self.labels = th["labels"]
        self.role = {s["role"]: s["label"] for s in t["sections"]}
        mcq = t["assessment"]["mcq"]
        self.options = mcq["option_labels"] if mcq["enabled"] else []
        self.option_display = dict(zip(mcq["option_labels"], mcq["option_display_labels"] or mcq["option_labels"]))
        self.callout_by_id = {c["id"]: c for c in t["callouts"]}
        self.title, self.subtitle = b["identity"]["title"], b["identity"]["subtitle"]
        self.credits = ", ".join(a["credit_line"] for a in b["identity"]["authors"])
        self.audience = b["audience"]["display_line"]
        self.preset = json.loads((REPO / "harness" / "presets" / f"{th['preset']}.json").read_text(encoding="utf-8"))
        self.figures_dir = cfg.path("figures")
        man = figures.load(cfg)
        self.manifest = {f["id"]: f for f in man["figures"]} if man is not None else None
        self.fig_numbers = {k: label for k, (_, label) in figures.numbering(cfg).items()} if man is not None else {}
        cover = th["cover"]["asset"]
        self.cover = (cfg["project"] / cover).resolve() if cover else None

    def style_of(self, key):
        """theme.callouts entry by callout id or section id -> (fill, label colour, layout, border)."""
        e = self.cfg["theme"]["callouts"].get(key)
        if not e:
            return DEFAULT_BOX, self.primary, "box", None
        return e["fill"].lstrip("#"), e["label_colour"].lstrip("#"), e["layout"], e.get("border")

    def question(self, n):
        """Display of the question marker `Q<n>` (theme.labels.question_prefix, Arabic contract P1)."""
        return f"{self.labels['question_prefix']}{n}"

    def objective(self, n):
        return f"{self.labels['objective_prefix']}{n}"


# ---------- the one element factory (EXT-LOC-2) ----------

class Factory:
    def __init__(self, d, theme):
        self.d, self.theme = d, theme
        if theme["direction"] != "ltr":
            raise BuildError("RTL rendering is STEP 10 (Arabic contract §3); this builder renders LTR only")

    def make_paragraph(self, container, text=None, style=None):
        return container.add_paragraph(text, style=style) if text is not None else (
            container.add_paragraph(style=style) if style else container.add_paragraph())

    def make_run(self, p, text, *, bold=False, italic=False, colour=None, size=None, sub=False, sup=False):
        r = p.add_run(text)
        r.bold = bold or None
        r.italic = italic or None
        if sub:
            r.font.subscript = True
        elif sup:
            r.font.superscript = True
        if colour:
            r.font.color.rgb = self.d["RGBColor"].from_string(colour)
        if size:
            r.font.size = self.d["Pt"](size)
        return r

    def make_table(self, container, rows, cols):
        return container.add_table(rows=rows, cols=cols)


# ---------- low-level XML helpers ----------

class Xml:
    def __init__(self, d):
        self.d = d
        self.el, self.qn = d["OxmlElement"], d["qn"]

    def add_math(self, p, tex, display=False):
        """Append one equation to a paragraph as OMML, the object Word's own equation editor produces."""
        from lxml import etree
        frag = etree.fromstring(mathml.omml_xml(tex, display))
        p._p.append(frag)
        return frag

    def shade(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = self.el("w:shd")
        shd.set(self.qn("w:val"), "clear")
        shd.set(self.qn("w:color"), "auto")
        shd.set(self.qn("w:fill"), fill)
        tcPr.append(shd)

    def cell_margins(self, cell, cm=0.3):
        tcPr = cell._tc.get_or_add_tcPr()
        mar = self.el("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            e = self.el(f"w:{side}")
            e.set(self.qn("w:w"), str(int(cm * 567)))
            e.set(self.qn("w:type"), "dxa")
            mar.append(e)
        tcPr.append(mar)

    def table_borders(self, table, spec):
        """spec: side -> (size_eighths, colour) or None for no border; sides incl. insideH/insideV."""
        b = self.el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = self.el(f"w:{side}")
            s = spec.get(side)
            if s:
                e.set(self.qn("w:val"), "single")
                e.set(self.qn("w:sz"), str(s[0]))
                e.set(self.qn("w:color"), s[1])
            else:
                e.set(self.qn("w:val"), "nil")
            b.append(e)
        table._tbl.tblPr.append(b)

    def cant_split(self, row):
        row._tr.get_or_add_trPr().append(self.el("w:cantSplit"))

    def para_border(self, p, side, size, colour, space=4):
        pPr = p._p.get_or_add_pPr()
        bdr = pPr.find(self.qn("w:pBdr"))
        if bdr is None:
            bdr = self.el("w:pBdr")
            pPr.append(bdr)
        e = self.el(f"w:{side}")
        e.set(self.qn("w:val"), "single")
        e.set(self.qn("w:sz"), str(size))
        e.set(self.qn("w:space"), str(space))
        e.set(self.qn("w:color"), colour)
        bdr.append(e)

    def add_field(self, p, instr, placeholder=""):
        r = p.add_run()
        for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
            if kind == "instr":
                t = self.el("w:instrText")
                t.set(self.qn("xml:space"), "preserve")
                t.text = text
                r._r.append(t)
            elif kind == "text":
                t = self.el("w:t")
                t.text = text
                r._r.append(t)
            else:
                f = self.el("w:fldChar")
                f.set(self.qn("w:fldCharType"), kind)
                r._r.append(f)
        return r

    def add_hyperlink(self, p, text, url, bold=False, italic=False):
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        h = self.el("w:hyperlink")
        h.set(self.qn("r:id"), rid)
        r = self.el("w:r")
        rPr = self.el("w:rPr")
        st = self.el("w:rStyle")
        st.set(self.qn("w:val"), "Hyperlink")
        rPr.append(st)
        if bold:
            rPr.append(self.el("w:b"))
        if italic:
            rPr.append(self.el("w:i"))
        r.append(rPr)
        t = self.el("w:t")
        t.text = text
        t.set(self.qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def page_numbering(self, section, fmt=None, start=None):
        pg = self.el("w:pgNumType")
        if fmt:
            pg.set(self.qn("w:fmt"), fmt)
        if start is not None:
            pg.set(self.qn("w:start"), str(start))
        section._sectPr.append(pg)

    def columns(self, section, n, space_cm=0.8):
        cols = section._sectPr.find(self.qn("w:cols"))
        if cols is None:
            cols = self.el("w:cols")
            section._sectPr.append(cols)
        cols.set(self.qn("w:num"), str(n))
        cols.set(self.qn("w:space"), str(int(space_cm * 567)))


# ---------- styles ----------

ALIGN = {"left": "LEFT", "justify": "JUSTIFY", "right": "RIGHT", "center": "CENTER"}


def set_font(d, style, name, size=None, bold=None, italic=None, colour=None):
    f = style.font
    f.name = name
    rf = style.element.get_or_add_rPr().find(d["qn"]("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(d["qn"](a), name)
    if size:
        f.size = d["Pt"](size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if colour:
        f.color.rgb = d["RGBColor"].from_string(colour)


def pstyle(d, doc, name, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        s = doc.styles.add_style(name, d["WD_STYLE_TYPE"].PARAGRAPH)
        s.base_style = doc.styles[base]
        return s


def setup_styles(d, doc, bk):
    Pt, Cm, A = d["Pt"], d["Cm"], d["WD_ALIGN_PARAGRAPH"]
    th = bk.cfg["theme"]
    body_align = getattr(A, ALIGN[th["alignment"]["body"]])
    head_align = getattr(A, ALIGN[th["alignment"]["headings"]])
    st = doc.styles
    normal = st["Normal"]
    set_font(d, normal, bk.serif, 10.5, colour=bk.ink)
    lang = d["OxmlElement"]("w:lang")
    lang.set(d["qn"]("w:val"), th["lang_tag"])
    normal.element.get_or_add_rPr().append(lang)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = d["WD_LINE_SPACING"].MULTIPLE
    pf.line_spacing = 1.18
    pf.alignment = body_align
    pf.widow_control = True

    h1 = st["Heading 1"]
    set_font(d, h1, bk.serif_head, 25, bold=True, italic=False, colour=bk.primary)
    h1.paragraph_format.space_before = Pt(64)
    h1.paragraph_format.space_after = Pt(22)
    h1.paragraph_format.alignment = head_align
    h1.paragraph_format.line_spacing = 1.0
    h1.paragraph_format.keep_with_next = True

    h2 = st["Heading 2"]
    set_font(d, h2, bk.sans, 12.5, bold=True, italic=False, colour=bk.primary)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.alignment = head_align
    h2.paragraph_format.keep_with_next = True

    s = pstyle(d, doc, "Heading 2 Unlisted", "Normal")
    set_font(d, s, bk.sans, 12.5, bold=True, colour=bk.primary)
    s.paragraph_format.space_before = Pt(16)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.alignment = head_align
    s.paragraph_format.keep_with_next = True

    s = pstyle(d, doc, "Title Heading", "Normal")
    set_font(d, s, bk.serif_head, 25, bold=True, colour=bk.primary)
    s.paragraph_format.space_before = Pt(64)
    s.paragraph_format.space_after = Pt(22)
    s.paragraph_format.alignment = head_align
    s.paragraph_format.line_spacing = 1.0

    h3 = st["Heading 3"]
    set_font(d, h3, bk.sans, 11, bold=True, italic=False, colour=bk.ink)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        rpr = st[name].element.get_or_add_rPr()
        for c in rpr.findall(d["qn"]("w:color")):   # drop the theme colour so the RGB wins
            for a in ("w:themeColor", "w:themeShade"):
                if c.get(d["qn"](a)):
                    del c.attrib[d["qn"](a)]

    s = pstyle(d, doc, "Box Label")
    set_font(d, s, bk.sans, 8.5, bold=True, colour=bk.primary)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.alignment = head_align
    s.paragraph_format.keep_with_next = True
    s.font.all_caps = True

    s = pstyle(d, doc, "Box Text")
    set_font(d, s, bk.serif, 9.8, colour=bk.ink)
    s.paragraph_format.space_after = Pt(4)

    s = pstyle(d, doc, "Caption Text")
    set_font(d, s, bk.sans, 8.5, colour=bk.muted)
    s.paragraph_format.alignment = head_align
    s.paragraph_format.space_after = Pt(12)

    s = pstyle(d, doc, "Figure")
    s.paragraph_format.alignment = A.CENTER
    s.paragraph_format.space_before = Pt(8)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.keep_with_next = True

    s = pstyle(d, doc, "Question")
    s.paragraph_format.space_before = Pt(8)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = head_align

    lists = th["lists"]
    s = pstyle(d, doc, "Option")
    set_font(d, s, bk.serif, 10, colour=bk.ink)
    s.paragraph_format.left_indent = Cm(lists["indent"])
    s.paragraph_format.first_line_indent = Cm(-lists["hanging"])
    s.paragraph_format.space_after = Pt(1)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = head_align

    s = pstyle(d, doc, "Answer")
    set_font(d, s, bk.serif, 9.8, colour=bk.ink)
    s.paragraph_format.space_after = Pt(6)

    s = pstyle(d, doc, "Reference")
    set_font(d, s, bk.serif, 8.8, colour=bk.ink)
    s.paragraph_format.left_indent = Cm(0.8)
    s.paragraph_format.first_line_indent = Cm(-0.8)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.alignment = head_align
    s.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))

    s = pstyle(d, doc, "Glossary Entry")
    set_font(d, s, bk.serif, 9.3, colour=bk.ink)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.alignment = head_align

    s = pstyle(d, doc, "Table Text")
    set_font(d, s, bk.sans, 8.8, colour=bk.ink)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.alignment = head_align
    s.paragraph_format.line_spacing = 1.05

    for name in ("Header", "Footer"):
        tabs = st[name].element.pPr.find(d["qn"]("w:tabs")) if st[name].element.pPr is not None else None
        if tabs is not None:
            tabs.getparent().remove(tabs)

    s = pstyle(d, doc, "Running Head")
    set_font(d, s, bk.sans, 8, colour=bk.muted)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.alignment = head_align

    for name in ("List Bullet", "List Number"):
        st[name].paragraph_format.space_after = Pt(3)
        st[name].paragraph_format.alignment = head_align

    hl = st["Hyperlink"] if "Hyperlink" in [x.name for x in st] else st.add_style("Hyperlink", d["WD_STYLE_TYPE"].CHARACTER)
    hl.font.color.rgb = d["RGBColor"].from_string(LINK)
    hl.font.underline = False


# ---------- renderer ----------

class Renderer:
    def __init__(self, d, doc, bk):
        self.d, self.doc, self.bk = d, doc, bk
        self.f, self.x = Factory(d, bk.cfg["theme"]), Xml(d)
        self.text_w = d["Cm"](bk.cfg["theme"]["layout"]["text_width"])

    # inline markdown
    def inline(self, p, text, bold=False, italic=False, colour=None, size=None, links=True):
        math_on = (self.bk.cfg["template"].get("math") or {}).get("enabled")
        for is_math, frag in (mathml.split_inline(text) if math_on else [(False, text)]):
            if is_math:
                self.x.add_math(p, frag)
                continue
            for sp in blocks.inline(frag, links):
                b, i = bold or sp["bold"], italic or sp["italic"]
                if sp["url"]:
                    self.x.add_hyperlink(p, sp["text"], sp["url"], b, i)
                else:
                    self.f.make_run(p, sp["text"], bold=b, italic=i, colour=colour, size=size,
                                    sub=sp["sub"], sup=sp["sup"])

    def reference_runs(self, p, text):
        """References: DOIs become links; the rest is inline markdown."""
        m = re.search(r"DOI:\s*" + DOI.pattern, text)
        if not m:
            self.inline(p, text)
            return
        self.inline(p, text[: m.start(1)])
        self.x.add_hyperlink(p, m.group(1), "https://doi.org/" + m.group(1))
        self.inline(p, text[m.end(1):])

    def box(self, container, label, fill, colour, border=None):
        t = self.f.make_table(container, 1, 1)
        t.alignment = self.d["WD_TABLE_ALIGNMENT"].CENTER
        t.autofit = False
        c = t.cell(0, 0)
        c.width = self.text_w
        self.x.shade(c, fill)
        self.x.cell_margins(c, 0.32)
        spec = {s: border for s in ("top", "left", "bottom", "right")} if border else {"top": (12, colour)}
        self.x.table_borders(t, spec)
        lp = c.paragraphs[0]
        lp.style = self.doc.styles["Box Label"]
        if label is not None:
            self.f.make_run(lp, label, bold=True, colour=colour)
        return t, c

    def after_table(self, container):
        p = self.f.make_paragraph(container)
        p.paragraph_format.space_after = self.d["Pt"](4)
        p.paragraph_format.line_spacing = self.d["Pt"](6)

    def grid(self, container, label, items, fill, colour):
        rows = 1 + math.ceil(len(items) / 2)
        t = self.f.make_table(container, rows, 2)
        t.alignment = self.d["WD_TABLE_ALIGNMENT"].CENTER
        t.autofit = False
        top = t.cell(0, 0).merge(t.cell(0, 1))
        for row in t.rows:
            self.x.cant_split(row)
            for c in row.cells:
                c.width = self.text_w // 2
                self.x.shade(c, fill)
                self.x.cell_margins(c, 0.3)
        self.x.table_borders(t, {"top": (12, colour), "insideH": (24, WHITE), "insideV": (24, WHITE)})
        lp = top.paragraphs[0]
        lp.style = self.doc.styles["Box Label"]
        self.f.make_run(lp, label, bold=True, colour=colour)
        A = self.d["WD_ALIGN_PARAGRAPH"]
        for k, (name, text) in enumerate(items):
            c = t.cell(1 + k // 2, k % 2)
            p = c.paragraphs[0]
            p.style = self.doc.styles["Box Text"]
            p.paragraph_format.alignment = A.LEFT
            self.f.make_run(p, name, bold=True, colour=colour)
            p2 = self.f.make_paragraph(c, style="Box Text")
            p2.paragraph_format.alignment = A.LEFT
            self.inline(p2, text)
        self.after_table(container)

    def callout(self, container, key, text, paras=None):
        """`key` is the callout ID from the parsed chapter (None: a quote block that is no configured callout).
        `paras` keeps the author's paragraph breaks inside the box; `text` is the whole body for a split layout."""
        label = self.bk.callout_by_id[key]["label"] if key else None
        fill, colour, layout, _ = self.bk.style_of(key)
        t, c = self.box(container, label, fill, colour)
        self.x.cant_split(t.rows[0])
        chunks = list(paras) if paras else [text]
        parts = self.bk.callout_by_id[key]["parts"] if key else []
        if layout == "split" and parts:
            rx = r"\s+".join(rf"{re.escape(pt)}:\s*(.+?)" for pt in parts[:-1]) + rf"\s+{re.escape(parts[-1])}:\s*(.+)"
            m = re.match(rx, text)
            if m:
                chunks = list(zip(parts, m.groups()))
        math_on = (self.bk.cfg["template"].get("math") or {}).get("enabled")
        for ch in chunks:
            if isinstance(ch, dict):   # a table the author put inside the box
                self.md_table(c, ch["rows"])
                continue
            p = self.f.make_paragraph(c, style="Box Text")
            if isinstance(ch, tuple):
                self.f.make_run(p, ch[0] + "  ", bold=True, colour=colour)
                self.inline(p, ch[1])
                continue
            tex = mathml.display(ch) if math_on else None
            if tex is not None:      # a line in the box that is only `$$...$$` is a display equation
                p.alignment = self.d["WD_ALIGN_PARAGRAPH"].CENTER
                self.x.add_math(p, tex, display=True)
            else:
                self.inline(p, ch)
        self.after_table(container)

    def md_table(self, container, cells):
        n = len(cells[0])
        t = self.f.make_table(container, len(cells), n)
        t.alignment = self.d["WD_TABLE_ALIGNMENT"].CENTER
        self.x.table_borders(t, {"top": (8, self.bk.primary), "bottom": (8, self.bk.primary), "insideH": (4, GRID)})
        for i, r in enumerate(cells):
            self.x.cant_split(t.rows[i])
            for j in range(n):
                c = t.cell(i, j)
                self.x.cell_margins(c, 0.15)
                p = c.paragraphs[0]
                p.style = self.doc.styles["Table Text"]
                txt = r[j] if j < len(r) else ""
                if i == 0:
                    self.x.shade(c, self.bk.primary)
                    self.inline(p, txt, bold=True, colour=WHITE)
                else:
                    if i % 2 == 0:
                        self.x.shade(c, STRIPE)
                    self.inline(p, txt)
        t.rows[0]._tr.get_or_add_trPr().append(self.d["OxmlElement"]("w:tblHeader"))   # repeat header row
        self.after_table(container)

    def figure_file(self, source, link):
        if link.startswith("fig:"):
            fig = (self.bk.manifest or {}).get(link[4:])
            if fig is None:
                raise BuildError(f"figure {link}: not in the figure manifest")
            src = figures.out_paths(self.bk.cfg, fig)[1]
            if not src.is_file():
                raise BuildError(f"figure {link}: {src} not rendered")
            return src
        src = assemble.resolve_asset(source, link, self.bk.cfg)
        if self.bk.manifest is not None:
            if self.bk.cover and src.resolve() == self.bk.cover:
                return src
            raise BuildError(f"image {link} is placed without the figure manifest (use fig:<id>)")
        if src.suffix == ".svg":   # legacy book: the PNG mirror beside the SVG sources
            src = self.bk.figures_dir / "png" / (src.stem + ".png")
        if not src.is_file():
            raise BuildError(f"figure {link}: {src} missing")
        return src

    def figure(self, container, src, caption, fig=None):
        """`fig`: the manifest entry (caption, alt and number from the manifest); None: a legacy figure."""
        from PIL import Image
        Cm = self.d["Cm"]
        with Image.open(src) as im:
            w_px, h_px = im.size
        width = Cm(figures.placed_width_cm(w_px, h_px, self.bk.cfg["theme"]["layout"]["text_width"]))
        p = self.f.make_paragraph(container, style="Figure")
        shape = p.add_run().add_picture(str(src), width=width)
        if fig is not None:
            shape._inline.docPr.set("descr", fig["alt"])
            cp = self.f.make_paragraph(container, style="Caption Text")
            self.f.make_run(cp, self.bk.fig_numbers[fig["id"]] + "  ", bold=True, colour=self.bk.primary)
            self.inline(cp, fig["caption"])
            if fig["kind"] == annotated.KIND:   # the numbered key under the caption (Task 9b.4)
                kp = self.f.make_paragraph(container, style="Caption Text")
                for k, (n, label) in enumerate(annotated.key(self.bk.cfg, fig)):
                    self.f.make_run(kp, ("   " if k else "") + f"{n} ", bold=True, colour=self.bk.primary)
                    self.inline(kp, label, colour=self.bk.ink)
            credit = self.f.make_paragraph(container, style="Caption Text")
            self.f.make_run(credit, f"{fig['credit']} ({fig['licence']})", colour=self.bk.muted, size=8)
            return
        if caption:
            cp = self.f.make_paragraph(container, style="Caption Text")
            m = re.match(rf"({re.escape(self.bk.labels['figure'])} \d+\.\d+)\s*[—-]\s*(.+)", caption)
            if m:
                self.f.make_run(cp, m.group(1) + "  ", bold=True, colour=self.bk.primary)
                self.inline(cp, m.group(2))
            else:
                self.inline(cp, caption)

    def markdown(self, text, kind, source):
        """kind: 'chapter' or 'glossary'; `source` is the file the text came from (image links are relative to it).
        The blocks come from the shared grammar (blocks.parse), which the HTML writer reads too."""
        d, doc, bk, f = self.d, self.doc, self.bk, self.f
        Pt, Cm, A = d["Pt"], d["Cm"], d["WD_ALIGN_PARAGRAPH"]
        container, role = doc, None
        resolved = {}

        def is_cover(link):
            resolved[link] = self.figure_file(source, link)
            return bool(bk.cover) and resolved[link].resolve() == bk.cover

        for b in blocks.parse(text, bk.cfg, kind, is_cover):
            t = b["t"]
            if t == "para":
                style = "Answer" if role == "answers" else None
                if container is not doc:
                    style = "Box Text"
                tex = mathml.display(b["text"]) if (bk.cfg["template"].get("math") or {}).get("enabled") else None
                if tex is not None:      # a paragraph that is only `$$...$$` is a display equation
                    mp = f.make_paragraph(container, style=(bk.cfg["template"]["math"].get("display_style") or style))
                    mp.alignment = self.d["WD_ALIGN_PARAGRAPH"].CENTER
                    self.x.add_math(mp, tex, display=True)
                else:
                    self.inline(f.make_paragraph(container, style=style), b["text"])
            elif t == "section":
                if container is not doc:
                    self.after_table(doc)
                container, role = doc, b["role"]
                if b["boxed"]:
                    fill, colour, _, border = bk.style_of(b["id"])
                    bd = (border["size_eighths_pt"], border["colour"].lstrip("#")) if border else None
                    _, container = self.box(doc, b["label"], fill, colour, bd)
                else:
                    f.make_paragraph(doc, b["label"], style="Heading 2 Unlisted" if b["unlisted"] else "Heading 2")
            elif t == "h3":   # inline markup rendered, like every other block (step9b fix S9b-01)
                self.inline(f.make_paragraph(container, style="Heading 3"), b["text"])
            elif t == "grid":
                fill, colour, _, _ = bk.style_of(b["key"])
                self.grid(container, b["label"], b["items"], fill, colour)
            elif t == "callout":
                self.callout(container, b["key"], b["text"], b.get("paras"))
            elif t == "table":
                self.md_table(container, b["rows"])
            elif t == "image":
                src = resolved[b["link"]]
                if b["link"].startswith("fig:"):
                    self.figure(container, src, None, bk.manifest[b["link"][4:]])
                else:
                    self.figure(container, src, b["caption"] or b["alt"])
            elif t == "question":
                p = f.make_paragraph(container, style="Question")
                f.make_run(p, bk.question(b["num"]) + "  ", bold=True, colour=bk.primary)
                self.inline(p, b["stem"])
                if b["los"]:
                    f.make_run(p, "  " + ", ".join(bk.objective(x) for x in b["los"]), colour=bk.muted, size=8)
                for k, (label, body) in enumerate(b["options"]):
                    op = f.make_paragraph(container, style="Option")
                    f.make_run(op, bk.option_display[label] + "\t", bold=True, colour=bk.primary)
                    self.inline(op, body)
                    op.paragraph_format.tab_stops.add_tab_stop(Cm(bk.cfg["theme"]["lists"]["indent"]))
                    if k == len(b["options"]) - 1:
                        op.paragraph_format.keep_with_next = False
                        op.paragraph_format.space_after = Pt(6)
            elif t == "bullet":
                p = f.make_paragraph(container, style="List Bullet")
                if container is not doc:
                    p.paragraph_format.space_after = Pt(2)
                self.inline(p, b["text"])
            elif t == "numbered":
                num, body = b["num"], b["body"]
                if role == "references":
                    p = f.make_paragraph(container, style="Reference")
                    f.make_run(p, num + ".\t", colour=bk.muted)
                    self.reference_runs(p, body)
                elif role == "objectives":
                    lo = LO_ITEM.match(body)
                    p = f.make_paragraph(container, style="Box Text")
                    p.paragraph_format.left_indent = Cm(1.1)
                    p.paragraph_format.first_line_indent = Cm(-1.1)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.1))
                    p.paragraph_format.alignment = A.LEFT
                    f.make_run(p, (bk.objective(lo.group(1)[2:]) if lo else num + ".") + "\t", bold=True, colour=bk.accent, size=8.5)
                    self.inline(p, lo.group(2) if lo else body)
                else:
                    p = f.make_paragraph(container, style="Box Text" if container is not doc else None)
                    p.paragraph_format.left_indent = Cm(0.7)
                    p.paragraph_format.first_line_indent = Cm(-0.7)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
                    f.make_run(p, num + ".\t", bold=True, colour=bk.primary)
                    self.inline(p, body)
            elif t == "glossary":
                self.inline(f.make_paragraph(container, style="Glossary Entry"), b["text"])
            elif t == "answer":
                p = f.make_paragraph(container, style="Answer")
                key_m = re.fullmatch(r"Q(\d+)\. (\S+)", b["head"])
                head = (f"{bk.question(key_m.group(1))}. {bk.option_display.get(key_m.group(2), key_m.group(2))}"
                        if key_m else b["head"])
                f.make_run(p, head + " ", bold=True, colour=bk.primary)
                self.inline(p, b["text"])
        if container is not doc:
            self.after_table(doc)


# ---------- document assembly ----------

class Pages:
    def __init__(self, d, doc, bk, R):
        self.d, self.doc, self.bk, self.R, self.x = d, doc, bk, R, R.x
        self.page = bk.cfg["theme"]["page"]
        self.preset_page = bk.preset["page"]

    def new_section(self, start_type=None):
        s = self.doc.add_section(start_type if start_type is not None else self.d["WD_SECTION"].NEW_PAGE)
        for pg in s._sectPr.findall(self.d["qn"]("w:pgNumType")):   # copied from the previous section; would restart numbering
            s._sectPr.remove(pg)
        for c in s._sectPr.findall(self.d["qn"]("w:cols")):
            s._sectPr.remove(c)
        return s

    def page_setup(self, section):
        Cm, pg = self.d["Cm"], self.page
        if pg["size"] != self.preset_page["size"]:
            raise BuildError(f"page size {pg['size']} is not the {self.bk.preset['id']} preset's ({self.preset_page['size']})")
        section.page_width, section.page_height = Cm(self.preset_page["width_cm"]), Cm(self.preset_page["height_cm"])
        section.left_margin, section.right_margin = Cm(pg["margins"]["left"]), Cm(pg["margins"]["right"])
        section.top_margin, section.bottom_margin = Cm(pg["margins"]["top"]), Cm(pg["margins"]["bottom"])
        section.header_distance, section.footer_distance = Cm(pg["header_distance"]), Cm(pg["footer_distance"])

    def running_head(self, section, left, right):
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        section.first_page_header.is_linked_to_previous = False
        section.first_page_footer.is_linked_to_previous = False
        section.different_first_page_header_footer = True
        self.fmt_head(section.header.paragraphs[0], left, right)
        for foot in (section.footer, section.first_page_footer):
            f = foot.paragraphs[0]
            f.alignment = self.d["WD_ALIGN_PARAGRAPH"].CENTER
            r = self.x.add_field(f, "PAGE", "1")
            r.font.name = self.bk.sans
            r.font.size = self.d["Pt"](8.5)
            r.font.color.rgb = self.d["RGBColor"].from_string(self.bk.muted)
        section.first_page_header.paragraphs[0].text = ""

    def fmt_head(self, p, left, right):
        p.paragraph_format.tab_stops.add_tab_stop(self.R.text_w, alignment=2)  # right
        for txt, bold in ((left, False), ("\t", False), (right, True)):
            r = p.add_run(txt)
            r.font.name = self.bk.sans
            r.font.size = self.d["Pt"](8)
            r.bold = bold or None
            r.font.color.rgb = self.d["RGBColor"].from_string(self.bk.muted)
        self.x.para_border(p, "bottom", 4, GRID, 3)

    @staticmethod
    def no_head(section):
        for hf in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            hf.is_linked_to_previous = False
            for p in hf.paragraphs:
                p.text = ""

    def heading1(self, number, title):
        p = self.R.f.make_paragraph(self.doc, style="Heading 1")
        if number:
            r = p.add_run(f"{number}")
            r.font.color.rgb = self.d["RGBColor"].from_string(self.bk.accent)
            p.add_run(GAP)
        if blocks.plain(title) == title:   # plain title: one bare run, as before (keeps existing DOCX output byte-identical)
            p.add_run(title)
        else:   # inline markup, the same grammar as the HTML writer (step9b fix S9b-01)
            self.R.inline(p, title, links=False)
        self.x.para_border(p, "bottom", 18, RULE, 10)
        return p

    def _layout_run(self, p, text, spec):
        Pt = self.d["Pt"]
        r = p.add_run(text)
        r.font.name, r.font.size = self.bk.fonts[spec["font"]], Pt(spec["size_pt"])
        if spec.get("bold"):
            r.bold = True
        if spec.get("italic"):
            r.italic = True
        if spec.get("colour"):
            r.font.color.rgb = self.d["RGBColor"].from_string(self.bk.pal[spec["colour"]].lstrip("#"))
        return r

    def title_pages(self):
        d, doc, bk, x = self.d, self.doc, self.bk, self.x
        Pt, Cm, A = d["Pt"], d["Cm"], d["WD_ALIGN_PARAGRAPH"]
        s = doc.sections[0]
        self.page_setup(s)
        if bk.cover:   # cover: zero-margin section with a full-bleed image
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(0)
            s.header_distance = s.footer_distance = Cm(0)
            self.no_head(s)
            p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(str(bk.cover), width=Cm(self.preset_page["width_cm"]),
                                    height=Cm(self.preset_page["height_cm"] - 0.01))
            self.new_section()   # move the section break into the cover paragraph so no blank page follows
            brk = doc.paragraphs[-1]._p
            sect = brk.pPr.find(d["qn"]("w:sectPr")) if brk.pPr is not None else None
            if sect is not None:
                p._p.get_or_add_pPr().append(sect)
                brk.getparent().remove(brk)
            s = doc.sections[-1]
            self.page_setup(s)
        self.no_head(s)
        x.page_numbering(s, "lowerRoman", 1)

        values = {"title": bk.title, "subtitle": bk.subtitle, "credits": bk.credits, "audience_line": bk.audience}
        notices = bk.cfg["theme"]["title_page"]["notices"]
        for spec in bk.preset["title_page_layout"]:
            if spec["field"] == "notices":
                if not notices:
                    continue
                p = doc.add_paragraph()
                if spec.get("new_page"):
                    p.add_run().add_break(d["WD_BREAK"].PAGE)
                p.paragraph_format.space_before = Pt(spec["space_before_pt"])
                for k, text in enumerate(notices):
                    q = doc.add_paragraph() if k else p
                    q.alignment = A.LEFT
                    self._layout_run(q, text, spec)
                    q.paragraph_format.space_after = Pt(spec["space_after_pt"])
                continue
            text = values[spec["field"]]
            if not text:
                continue
            p = doc.add_paragraph()
            if spec.get("space_before_pt"):
                p.paragraph_format.space_before = Pt(spec["space_before_pt"])
            p.alignment = A.LEFT
            self._layout_run(p, text, spec)
            rule = spec.get("rule_below")
            if rule:
                x.para_border(p, "bottom", rule["size_eighths_pt"], rule["colour"], rule["space_pt"])

        p = doc.add_paragraph(bk.labels["contents"], style="Title Heading")
        p.paragraph_format.page_break_before = True
        x.para_border(p, "bottom", 18, RULE, 10)
        tp = doc.add_paragraph()
        levels = bk.cfg["theme"]["toc"]["levels"]
        x.add_field(tp, f'TOC \\o "1-{levels}" \\h \\z \\u', bk.labels["update_toc_instruction"])


def chapter_heading(text, cfg):
    m = re.search(cfg["template"]["chapter_heading_pattern"], text, re.M)
    if not m:
        raise BuildError("chapter heading does not match template.chapter_heading_pattern")
    return int(m.group("num")), m.group("title")


def docx_path(cfg):
    return cfg["project"] / "build" / f"{cfg['theme']['output']['basename']}.docx"


def build(cfg, out=None):
    """-> the written .docx: `out`, default build/<basename>.docx (only `run build` writes there)."""
    d = _docx()
    bk = Book(cfg)
    doc = d["Document"]()
    setup_styles(d, doc, bk)
    R = Renderer(d, doc, bk)
    pages = Pages(d, doc, bk, R)
    x, Pt, Cm, A = R.x, d["Pt"], d["Cm"], d["WD_ALIGN_PARAGRAPH"]
    pages.title_pages()

    front_file = cfg.book_file("front_matter")
    how = bk.role.get("how_to_use")
    if front_file is not None and how:
        front = front_file.read_text(encoding="utf-8")
        marker = f"\n## {how}"
        if marker in front:
            s = pages.new_section()
            pages.page_setup(s)
            x.page_numbering(s, "lowerRoman")
            pages.running_head(s, bk.title, how)
            p = R.f.make_paragraph(doc, how, style="Heading 1")
            x.para_border(p, "bottom", 18, RULE, 10)
            R.markdown(front.split(marker, 1)[1], "chapter", front_file)

    chapters = []
    for c, path in cfg.chapters():
        text = path.read_text(encoding="utf-8")
        chapters.append((c, path, text, *chapter_heading(text, cfg)))
    parts = {p["id"]: p for p in cfg["chapter_plan"].get("parts", [])}
    members = {}
    for c, _, _, n, _ in chapters:
        if c.get("part_id"):
            members.setdefault(c["part_id"], []).append(n)
    if parts and "chapters" not in bk.labels:
        raise BuildError("theme.labels.chapters is required when chapter-plan.json has parts")
    first, current = True, None
    for c, path, text, n, title in chapters:
        pid = c.get("part_id")
        if pid and pid != current:
            current = pid
            part = parts[pid]
            s = pages.new_section()
            pages.page_setup(s)
            pages.no_head(s)
            if first:
                x.page_numbering(s, "decimal", 1)
                first = False
            p = R.f.make_paragraph(doc, style="Heading 1")
            p.paragraph_format.space_before = Pt(200)
            r = p.add_run(part["label"])
            r.font.color.rgb = d["RGBColor"].from_string(bk.accent)
            p.add_run(GAP + part["title"])
            x.para_border(p, "bottom", 18, RULE, 10)
            q = doc.add_paragraph()
            span = members[pid]
            r = q.add_run(f"{bk.labels['chapters']} {span[0]}–{span[-1]}" if len(span) > 1 else f"{bk.labels['chapter']} {span[0]}")
            r.font.name, r.font.size = bk.sans, Pt(9.5)
            r.font.color.rgb = d["RGBColor"].from_string(bk.muted)
            q.paragraph_format.space_after = Pt(14)
            for c2, _, _, k, t2 in chapters:
                if c2.get("part_id") != pid:
                    continue
                q = doc.add_paragraph()
                q.alignment = A.LEFT
                q.paragraph_format.tab_stops.add_tab_stop(Cm(1))
                r = q.add_run(f"{k}\t")
                r.bold, r.font.color.rgb = True, d["RGBColor"].from_string(bk.accent)
                r = q.add_run(t2)
                r.font.name, r.font.size = bk.serif_head, Pt(13)
        s = pages.new_section(d["WD_SECTION"].NEW_PAGE)
        pages.page_setup(s)
        if first:
            x.page_numbering(s, "decimal", 1)
            first = False
        pages.running_head(s, bk.title, f"{bk.labels['chapter']} {n} · {blocks.plain(title)}")
        pages.heading1(n, title)
        R.markdown(text, "chapter", path)

    glossary = cfg.book_file("glossary") if cfg["template"]["glossary"]["enabled"] else None
    if glossary is not None:   # heading single column, entries in two columns
        s = pages.new_section()
        pages.page_setup(s)
        pages.running_head(s, bk.title, bk.labels["glossary"])
        pages.heading1(None, bk.labels["glossary"])
        s = pages.new_section(d["WD_SECTION"].CONTINUOUS)
        pages.page_setup(s)
        x.columns(s, 2)
        R.markdown(glossary.read_text(encoding="utf-8"), "glossary", glossary)

    settings = doc.settings.element
    settings.append(d["OxmlElement"]("w:autoHyphenation"))
    ul = d["OxmlElement"]("w:updateFields")
    ul.set(d["qn"]("w:val"), "true")
    settings.append(ul)

    core = doc.core_properties
    core.title = f"{bk.title}: {bk.subtitle}" if bk.subtitle else bk.title
    core.author = bk.credits
    core.language = cfg["theme"]["lang_tag"]
    core.comments = ""  # python-docx default is "generated by python-docx"
    core.last_modified_by = bk.credits
    out = out or docx_path(cfg)
    out.parent.mkdir(exist_ok=True)
    doc.save(out)
    return out


def _bgr(hexcol):
    return int(hexcol[4:6] + hexcol[2:4] + hexcol[0:2], 16)


def word_finish(cfg, docx, export_pdf=True):
    """Word COM backend (INV-38): hidden instance, update the contents, style it, embed fonts, save, export PDF
    (not when the theme's PDF engine is `html`: `export_pdf` False). -> (pdf path or None, Word's page count)."""
    import win32com.client as win32
    th = cfg["theme"]
    bk = Book(cfg)
    if th["build"]["backend"] != "word_com":
        raise BuildError(f"build backend {th['build']['backend']!r} is not supported")
    pdf = docx.with_suffix(".pdf")
    if export_pdf:
        try:
            open(pdf, "ab").close()
        except PermissionError:
            raise BuildError(f"{pdf.name} is open in another program; close it and re-run build")
    toc = [(bk.sans, 10.5, True, 8, bk.primary, 0), (bk.serif, 10, False, 0, bk.ink, 28)]   # preset TOC levels 1, 2
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(docx))
        for lvl, name in enumerate(th["build"]["word_com"]["toc_styles"][:len(toc)], 1):
            font, size, bold, before, colour, indent = toc[lvl - 1]
            st = doc.Styles(name)
            st.Font.Name = font
            st.Font.Size = size
            st.Font.Bold = bold
            st.Font.Color = _bgr(colour)
            st.ParagraphFormat.SpaceBefore = before
            st.ParagraphFormat.SpaceAfter = 2
            st.ParagraphFormat.LeftIndent = indent
        doc.TablesOfContents(1).Update()
        doc.Fields.Update()
        doc.TablesOfContents(1).Update()
        doc.EmbedTrueTypeFonts = bool(th["build"]["pdf"]["embed_fonts"])
        doc.SaveSubsetFonts = False
        doc.Save()
        if export_pdf:
            bookmarks = 1 if th["build"]["pdf"]["bookmarks"] == "headings" else 0
            # 17 = wdExportFormatPDF, 0 = print quality; CreateBookmarks: 1 = headings
            doc.ExportAsFixedFormat(str(pdf), 17, False, 0, 0, 1, 1, 0, True, True, bookmarks, True, True, False)
        pages = doc.ComputeStatistics(2)
        doc.Close(False)
    finally:
        word.Quit()
    return (pdf if export_pdf else None), pages


def main(project, argv):
    try:
        cfg = config.load(project)
        with tempfile.TemporaryDirectory() as t:
            out = build(cfg, pathlib.Path(t) / docx_path(cfg).name)
            print(f"built {out.name} (dry run, not kept)")
            if "--no-pdf" not in argv:
                pdf, pages = word_finish(cfg, out)
                print(f"built {pdf.name}: {pages} pages (dry run, not kept)")
    except (config.ConfigError, KeyError, FileNotFoundError) as e:
        print(f"ERROR CONFIG: {e}", file=sys.stderr)
        return 2
    except (BuildError, assemble.AssetOutside) as e:
        print(f"ERROR BUILD: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    from harness.gate import enforce
    PROJECT = enforce("build", sys.argv)   # first call: no stage work before the gates pass
    sys.exit(main(PROJECT, sys.argv))
