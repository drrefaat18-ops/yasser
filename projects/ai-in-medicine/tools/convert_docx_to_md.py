import os
import re
import sys
import zipfile
import docx
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

if __name__ == "__main__":
    import pathlib, sys  # gate shim (Task 7.3): no stage work before the gates pass
    sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[3]))  # repo root
    from harness.gate import enforce
    PROJECT = enforce("ingest", sys.argv)

sys.stdout.reconfigure(encoding='utf-8')

BOOK_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_PATH = os.path.join(BOOK_DIR, "original", "AI_in_Medicine- Assistant Prof Dr Shereen Elkholy.docx")
MD_PATH = os.path.join(BOOK_DIR, "original", "AI_in_Medicine- Assistant Prof Dr Shereen Elkholy.md")
IMAGES_DIR = os.path.join(BOOK_DIR, "images")

os.makedirs(IMAGES_DIR, exist_ok=True)

# 1. Extract all media files from docx zip
with zipfile.ZipFile(DOCX_PATH, 'r') as z:
    for name in z.namelist():
        if name.startswith('word/media/'):
            filename = os.path.basename(name)
            target_file = os.path.join(IMAGES_DIR, filename)
            with open(target_file, 'wb') as f:
                f.write(z.read(name))
print(f"Extracted media files to {IMAGES_DIR}")

# 2. Open document with python-docx
doc = docx.Document(DOCX_PATH)

def get_blip_targets(element):
    """Find all image references within an XML element and map to ../images/filename (MD_PATH is in original/)."""
    blips = element.xpath('.//a:blip')
    targets = []
    for b in blips:
        embed_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        if embed_id and embed_id in doc.part.rels:
            target_ref = doc.part.rels[embed_id].target_ref
            filename = os.path.basename(target_ref)
            targets.append(f"../images/{filename}")
    return targets

def split_paragraph_into_line_tokens(p):
    """Split a paragraph's runs across newlines to preserve per-line formatting."""
    lines = [[]]
    for r in p.runs:
        parts = r.text.split('\n')
        for idx, part in enumerate(parts):
            if idx > 0:
                lines.append([])
            if part:
                lines[-1].append({
                    'text': part,
                    'bold': bool(r.bold),
                    'italic': bool(r.italic),
                    'font_size': r.font.size.pt if r.font and r.font.size else None
                })
    return lines

def format_line_tokens(tokens):
    """Format token list into markdown string, returning (formatted_text, max_font_size, all_bold)."""
    if not tokens:
        return '', 0, False
    
    # Merge consecutive identical styling
    merged = []
    max_sz = 0
    all_bold = True
    for t in tokens:
        if t['font_size'] and t['font_size'] > max_sz:
            max_sz = t['font_size']
        if not t['bold'] and t['text'].strip():
            all_bold = False
            
        if merged and merged[-1]['bold'] == t['bold'] and merged[-1]['italic'] == t['italic']:
            merged[-1]['text'] += t['text']
            if t['font_size'] and (merged[-1]['font_size'] is None or t['font_size'] > merged[-1]['font_size']):
                merged[-1]['font_size'] = t['font_size']
        else:
            merged.append(dict(t))
            
    out = ''
    for item in merged:
        raw = item['text']
        b = item['bold']
        i = item['italic']
        if not b and not i:
            out += raw
            continue
        leading_ws = re.match(r'^\s*', raw).group()
        trailing_ws = re.search(r'\s*$', raw).group()
        core = raw[len(leading_ws):len(raw)-len(trailing_ws)]
        if not core:
            out += raw
            continue
        styled = core
        if i:
            styled = f'*{styled}*'
        if b:
            styled = f'**{styled}**'
        out += leading_ws + styled + trailing_ws
        
    return out, max_sz, all_bold

def process_paragraph(p, inside_blockquote=False, in_toc=False):
    """Convert a paragraph to markdown lines."""
    blip_targets = get_blip_targets(p._p)
    image_lines = [f"![Image]({t})" for t in blip_targets]
    
    raw_full = p.text.strip()
    if not raw_full:
        return image_lines
        
    line_tokens_list = split_paragraph_into_line_tokens(p)
    style_name = p.style.name
    
    lines = []
    for tokens in line_tokens_list:
        formatted_text, max_sz, all_bold = format_line_tokens(tokens)
        if not formatted_text.strip():
            continue
            
        clean_plain = re.sub(r'[*_`]', '', formatted_text).strip()
        raw_text = clean_plain
        
        # Check if TOC item
        if in_toc:
            m = re.match(r'^(Chapter\s+\d+:\s+.*?)\t+(\d+)$', clean_plain)
            if m:
                title, page = m.group(1), m.group(2)
                anchor = re.sub(r'[^\w\s-]', '', title.lower()).strip()
                anchor = re.sub(r'[\s]+', '-', anchor)
                lines.append(f"- [{title}](#{anchor}) *(p. {page})*")
                continue
            elif clean_plain == "Table of Contents":
                lines.append("## Table of Contents")
                continue
                
        # Determine heading level
        heading_level = 0
        if re.match(r'^\d+\.\d+\.\d+\s', clean_plain):
            heading_level = 3
        elif re.match(r'^\d+\.\d+\s', clean_plain):
            heading_level = 2
        elif clean_plain.startswith('Chapter Executive Overview'):
            heading_level = 2
        elif style_name == 'Heading 1':
            heading_level = 1
        elif style_name == 'Heading 2':
            heading_level = 2
        elif style_name == 'Heading 3':
            heading_level = 3
        elif clean_plain == "ARTIFICIAL INTELLIGENCE":
            heading_level = 1
        elif clean_plain == "Table of Contents":
            heading_level = 2
        elif clean_plain.startswith("References") and len(clean_plain) < 30:
            heading_level = 1
        elif re.match(r'^Chapter\s+\d+[:\s]', clean_plain, re.IGNORECASE):
            heading_level = 1
        elif re.match(r'^Chapter\s+\d+\s+Review', clean_plain, re.IGNORECASE) or "Review & Self-Assessment" in clean_plain:
            heading_level = 2
        elif clean_plain.startswith("Self-Assessment Quiz"):
            heading_level = 2
        elif max_sz >= 20 and len(clean_plain) < 100:
            heading_level = 1
        elif (max_sz >= 16 or (max_sz >= 14 and all_bold)) and len(clean_plain) < 120:
            if re.match(r'^\d+\.\d+\.\d+\s', clean_plain):
                heading_level = 3
            elif re.match(r'^\d+\.\d+\s', clean_plain):
                heading_level = 2
            elif max_sz >= 18:
                heading_level = 2
            elif clean_plain.startswith("Chapter Executive Overview"):
                heading_level = 2
                
        if heading_level > 0 and not inside_blockquote:
            prefix = "#" * heading_level
            h_text = clean_plain
            if re.match(r'^CHAPTER\s+(\d+):', h_text):
                h_text = re.sub(r'^CHAPTER\s+(\d+):', r'Chapter \1:', h_text)
            lines.append(f"{prefix} {h_text}")
        elif style_name == 'List Bullet' or raw_text.startswith('•') or raw_text.startswith('·') or raw_text.startswith('\uf0b7'):
            # Strip bullet character carefully without stripping markdown asterisks
            item_text = re.sub(r'^[•·\u2022\u25cf\u25cb\u25aa\u25ab\uf0b7\t ]+', '', formatted_text).strip()
            lines.append(f"- {item_text}")
        elif re.match(r'^\d+[\.\)]\s+', raw_text):
            lines.append(formatted_text.strip())
        else:
            lines.append(formatted_text.strip())
            
    lines.extend(image_lines)
    return lines

def process_table(tbl):
    """Convert a table to markdown block."""
    num_rows = len(tbl.rows)
    num_cols = len(tbl.columns)
    
    # 1x1 table (callout/figure card)
    if num_rows == 1 and num_cols == 1:
        cell = tbl.rows[0].cells[0]
        md_lines = []
        for p in cell.paragraphs:
            p_lines = process_paragraph(p, inside_blockquote=True)
            for l in p_lines:
                if l.strip():
                    md_lines.append(f"> {l}")
                else:
                    md_lines.append(">")
        return "\n".join(md_lines)
        
    # Multi-row / Multi-column table
    table_lines = []
    headers = []
    for c in tbl.rows[0].cells:
        cell_text = " ".join([format_line_tokens(split_paragraph_into_line_tokens(p)[0])[0].strip() 
                             for p in c.paragraphs if p.text.strip()])
        cell_text = cell_text.replace('\n', '<br>').replace('|', '\\|')
        headers.append(cell_text if cell_text else " ")
        
    table_lines.append("| " + " | ".join(headers) + " |")
    table_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
    
    for row in tbl.rows[1:]:
        row_cells = []
        for c in row.cells:
            cell_text = " ".join([format_line_tokens(split_paragraph_into_line_tokens(p)[0])[0].strip() 
                                 for p in c.paragraphs if p.text.strip()])
            cell_text = cell_text.replace('\n', '<br>').replace('|', '\\|')
            row_cells.append(cell_text if cell_text else " ")
        table_lines.append("| " + " | ".join(row_cells) + " |")
        
    return "\n".join(table_lines)

# 3. Main conversion loop
output_blocks = []
in_toc = False

for child in doc._body._body:
    if isinstance(child, CT_P):
        p = Paragraph(child, doc)
        raw_text = p.text.strip()
        
        # Detect Table of Contents section
        if raw_text == "Table of Contents":
            in_toc = True
        elif in_toc and (raw_text.startswith("Chapter 1:") and "\t" not in raw_text):
            in_toc = False
            
        p_lines = process_paragraph(p, in_toc=in_toc)
        for l in p_lines:
            if l.strip():
                output_blocks.append(l)
            else:
                output_blocks.append("")
    elif isinstance(child, CT_Tbl):
        t = Table(child, doc)
        tbl_md = process_table(t)
        if tbl_md.strip():
            output_blocks.append(tbl_md)

# 4. Clean up consecutive empty blocks and write output
cleaned = []
prev_blank = False
for b in output_blocks:
    is_blank = (b.strip() == "")
    if is_blank:
        if not prev_blank:
            cleaned.append("")
        prev_blank = True
    else:
        cleaned.append(b)
        prev_blank = False

md_content = "\n\n".join(cleaned)
md_content = re.sub(r'\n{3,}', '\n\n', md_content)

# Add missing Chapter 7 link in Table of Contents if not already present
if "[Chapter 6:" in md_content and "[Chapter 7:" not in md_content:
    md_content = re.sub(
        r'(\- \[Chapter 6:.*?\(p\. \d+\)\*)',
        r'\1\n- [Chapter 7: Ethics, Medicolegal Challenges & Future Horizons](#chapter-7-ethics-medicolegal-challenges--future-horizons)\n- [References](#references)',
        md_content
    )

if md_content.startswith('# ARTIFICIAL INTELLIGENCE\n\n# IN MEDICINE'):
    md_content = md_content.replace('# ARTIFICIAL INTELLIGENCE\n\n# IN MEDICINE', '# ARTIFICIAL INTELLIGENCE IN MEDICINE', 1)

with open(MD_PATH, 'w', encoding='utf-8') as f:
    f.write(md_content)

print(f"Successfully generated {MD_PATH}")
print(f"Total blocks: {len(cleaned)}, File size: {os.path.getsize(MD_PATH)} bytes")
