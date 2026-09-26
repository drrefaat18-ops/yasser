"""Build the textbook as Word (.docx) and PDF from the rework Markdown sources.

Usage: python projects/ai-in-medicine/rework/tools/build_book.py [--no-pdf]
Writes AI_in_Health_Care_Interprofessional.docx (and .pdf via Microsoft Word) in deliverables/.
Figures: SVGs must be pre-rendered to rework/figures/png/ (see render_figures in README of this script's docstring):
  chrome --headless --force-device-scale-factor=4 --screenshot=png/X.png X.svg
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parents[2]
REWORK = ROOT / "rework"
FIG = REWORK / "figures"
OUT = ROOT / "deliverables" / "AI_in_Health_Care_Interprofessional.docx"

BOOK = "Artificial Intelligence in Health Care"
SUB = "An Interprofessional Introduction"
AUTHOR = "Assistant Prof. Dr. Shereen Elsaid Elkholy"

SERIF, SERIF_HEAD, SANS = "Sitka Text", "Sitka Heading", "Segoe UI"
INK, NAVY, GOLD, MUTED = "1D2A38", "12324F", "9A6B00", "5A6470"
TEXT_W = Cm(14)

# box label -> (fill, label colour)
BOXES = {
    "Medical Background in 60 Seconds": ("E8F0F7", "2B5D8A"),
    "Through Four Lenses": ("F1F3F5", NAVY),
    "Myth vs Evidence": ("FBF3E0", "7A5200"),
    "Safety Alert": ("FBEAEA", "A12B2B"),
    "Deeper Dive": ("F1F1EE", "4A5560"),
    "Local Context": ("E6F2E6", "2F6A36"),
    "Opening Case": ("EEF2F6", NAVY),
    "Learning Objectives": ("FFFFFF", NAVY),
}
PARTS = {
    1: ("Part I", "Foundations", "Chapters 1–5"),
    6: ("Part II", "AI Across the Care Pathway", "Chapters 6–9"),
    10: ("Part III", "Responsible AI", "Chapters 10–11"),
}
BOXED_SECTIONS = {"Opening Case", "Learning Objectives"}
UNLISTED = {"Key Takeaways", "Answers and Rationales"}  # kept out of the contents page


# ---------- low-level XML helpers ----------

def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(cell, cm=0.3):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(cm * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def table_borders(table, spec):
    """spec: dict side -> (size_eighths, colour) or None for no border; sides incl. insideH/insideV."""
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        s = spec.get(side)
        if s:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(s[0]))
            e.set(qn("w:color"), s[1])
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def para_border(p, side, size, colour, space=4):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), colour)
    bdr.append(e)


def add_field(p, instr, placeholder=""):
    r = p.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r._r.append(t)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            r._r.append(t)
        else:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
    return r


def add_hyperlink(p, text, url, bold=False, italic=False):
    part = p.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    st = OxmlElement("w:rStyle")
    st.set(qn("w:val"), "Hyperlink")
    rPr.append(st)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def page_numbering(section, fmt=None, start=None):
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    section._sectPr.append(pg)


def columns(section, n, space_cm=0.8):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_cm * 567)))


# ---------- styles ----------

def set_font(style, name, size=None, bold=None, italic=None, colour=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if colour:
        f.color.rgb = rgb(colour)


def pstyle(doc, name, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = doc.styles[base]
        return s


def setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    set_font(normal, SERIF, 10.5, colour=INK)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-GB")
    normal.element.get_or_add_rPr().append(lang)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.18
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    h1 = st["Heading 1"]
    set_font(h1, SERIF_HEAD, 25, bold=True, italic=False, colour=NAVY)
    h1.paragraph_format.space_before = Pt(64)
    h1.paragraph_format.space_after = Pt(22)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.line_spacing = 1.0
    h1.paragraph_format.keep_with_next = True

    h2 = st["Heading 2"]
    set_font(h2, SANS, 12.5, bold=True, italic=False, colour=NAVY)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.keep_with_next = True

    s = pstyle(doc, "Heading 2 Unlisted", "Normal")
    set_font(s, SANS, 12.5, bold=True, colour=NAVY)
    s.paragraph_format.space_before = Pt(16)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.keep_with_next = True

    s = pstyle(doc, "Title Heading", "Normal")
    set_font(s, SERIF_HEAD, 25, bold=True, colour=NAVY)
    s.paragraph_format.space_before = Pt(64)
    s.paragraph_format.space_after = Pt(22)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.line_spacing = 1.0

    h3 = st["Heading 3"]
    set_font(h3, SANS, 11, bold=True, italic=False, colour=INK)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        rpr = st[name].element.get_or_add_rPr()
        # drop the theme colour so our RGB wins
        for c in rpr.findall(qn("w:color")):
            for a in ("w:themeColor", "w:themeShade"):
                if c.get(qn(a)):
                    del c.attrib[qn(a)]

    s = pstyle(doc, "Box Label")
    set_font(s, SANS, 8.5, bold=True, colour=NAVY)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.keep_with_next = True
    s.font.all_caps = True

    s = pstyle(doc, "Box Text")
    set_font(s, SERIF, 9.8, colour=INK)
    s.paragraph_format.space_after = Pt(4)

    s = pstyle(doc, "Caption Text")
    set_font(s, SANS, 8.5, colour=MUTED)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.space_after = Pt(12)

    s = pstyle(doc, "Figure")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_before = Pt(8)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.keep_with_next = True

    s = pstyle(doc, "Question")
    s.paragraph_format.space_before = Pt(8)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    s = pstyle(doc, "Option")
    set_font(s, SERIF, 10, colour=INK)
    s.paragraph_format.left_indent = Cm(1.2)
    s.paragraph_format.first_line_indent = Cm(-0.6)
    s.paragraph_format.space_after = Pt(1)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    s = pstyle(doc, "Answer")
    set_font(s, SERIF, 9.8, colour=INK)
    s.paragraph_format.space_after = Pt(6)

    s = pstyle(doc, "Reference")
    set_font(s, SERIF, 8.8, colour=INK)
    s.paragraph_format.left_indent = Cm(0.8)
    s.paragraph_format.first_line_indent = Cm(-0.8)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))

    s = pstyle(doc, "Glossary Entry")
    set_font(s, SERIF, 9.3, colour=INK)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    s = pstyle(doc, "Table Text")
    set_font(s, SANS, 8.8, colour=INK)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.line_spacing = 1.05

    for name in ("Header", "Footer"):
        tabs = st[name].element.pPr.find(qn("w:tabs")) if st[name].element.pPr is not None else None
        if tabs is not None:
            tabs.getparent().remove(tabs)

    s = pstyle(doc, "Running Head")
    set_font(s, SANS, 8, colour=MUTED)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("List Bullet", "List Number"):
        st[name].paragraph_format.space_after = Pt(3)
        st[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    hl = st["Hyperlink"] if "Hyperlink" in [x.name for x in st] else st.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    hl.font.color.rgb = rgb("2B5D8A")
    hl.font.underline = False


# ---------- inline markdown ----------

INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*])|\[[^\]]+\]\([^)\s]+\))")
DOI = re.compile(r"(10\.\d{4,9}/[^\s]+[^\s.,;])")
URL = re.compile(r"(https?://[^\s)]+[^\s).,;])")


def inline(p, text, bold=False, italic=False, colour=None, size=None, links=True):
    """Add runs for inline markdown to paragraph p."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        b, i, t, url = bold, italic, tok, None
        if tok.startswith("***") and tok.endswith("***") and len(tok) > 6:
            b, i, t = True, True, tok[3:-3]
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            b, t = True, tok[2:-2]
        elif tok.startswith("[") and "](" in tok and tok.endswith(")"):
            t, url = tok[1:].split("](", 1)
            url = url[:-1]
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            i, t = True, tok[1:-1]
        if url and links:
            add_hyperlink(p, t, url, b, i)
            continue
        if links and URL.search(t):
            parts = URL.split(t)
            for k, seg in enumerate(parts):
                if not seg:
                    continue
                if k % 2:
                    add_hyperlink(p, seg, seg, b, i)
                else:
                    _run(p, seg, b, i, colour, size)
            continue
        _run(p, t, b, i, colour, size)


def _run(p, t, b, i, colour, size):
    r = p.add_run(t)
    r.bold = b or None
    r.italic = i or None
    if colour:
        r.font.color.rgb = rgb(colour)
    if size:
        r.font.size = Pt(size)
    return r


def reference_runs(p, text):
    """References: make DOIs clickable; keep the rest as inline markdown."""
    m = re.search(r"DOI:\s*" + DOI.pattern, text)
    if not m:
        inline(p, text)
        return
    inline(p, text[: m.start(1)])
    add_hyperlink(p, m.group(1), "https://doi.org/" + m.group(1))
    inline(p, text[m.end(1):])


# ---------- block renderers ----------

class Renderer:
    def __init__(self, doc):
        self.doc = doc

    def box(self, container, label, fill, colour, border=None):
        t = container.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        c = t.cell(0, 0)
        c.width = TEXT_W
        shade(c, fill)
        cell_margins(c, 0.32)
        spec = {s: border for s in ("top", "left", "bottom", "right")} if border else {}
        if not border:
            spec = {"top": (12, colour)}
        table_borders(t, spec)
        lp = c.paragraphs[0]
        lp.style = self.doc.styles["Box Label"]
        _run(lp, label, True, False, colour, None)
        return t, c

    def after_table(self, container):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(6)
        for r in p.runs:
            r.font.size = Pt(4)

    def lenses(self, container, items):
        fill, colour = BOXES["Through Four Lenses"]
        t = container.add_table(rows=3, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        top = t.cell(0, 0).merge(t.cell(0, 1))
        for row in t.rows:
            cant_split(row)
            for c in row.cells:
                c.width = TEXT_W // 2
                shade(c, fill)
                cell_margins(c, 0.3)
        table_borders(t, {"top": (12, colour), "insideH": (24, "FFFFFF"), "insideV": (24, "FFFFFF")})
        lp = top.paragraphs[0]
        lp.style = self.doc.styles["Box Label"]
        _run(lp, "Through Four Lenses", True, False, colour, None)
        for k, (prof, text) in enumerate(items[:4]):
            c = t.cell(1 + k // 2, k % 2)
            p = c.paragraphs[0]
            p.style = self.doc.styles["Box Text"]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _run(p, prof, True, False, colour, None)
            p2 = c.add_paragraph(style="Box Text")
            p2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(p2, text)
        self.after_table(container)

    def callout(self, container, label, text):
        fill, colour = BOXES.get(label, ("F1F1EE", NAVY))
        t, c = self.box(container, label, fill, colour)
        cant_split(t.rows[0])
        chunks = [text]
        if label == "Myth vs Evidence":
            m = re.match(r"Myth:\s*(.+?)\s+Evidence:\s*(.+)", text)
            if m:
                chunks = [("Myth", m.group(1)), ("Evidence", m.group(2))]
        for ch in chunks:
            p = c.add_paragraph(style="Box Text")
            if isinstance(ch, tuple):
                _run(p, ch[0] + "  ", True, False, colour, None)
                inline(p, ch[1])
            else:
                inline(p, ch)
        self.after_table(container)

    def md_table(self, container, rows):
        cells = [[x.strip() for x in r.strip().strip("|").split("|")] for r in rows]
        cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", x) for x in r)]
        n = len(cells[0])
        t = container.add_table(rows=len(cells), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t, {"top": (8, NAVY), "bottom": (8, NAVY), "insideH": (4, "C9D1D9")})
        for i, r in enumerate(cells):
            cant_split(t.rows[i])
            for j in range(n):
                c = t.cell(i, j)
                cell_margins(c, 0.15)
                p = c.paragraphs[0]
                p.style = self.doc.styles["Table Text"]
                txt = r[j] if j < len(r) else ""
                if i == 0:
                    shade(c, NAVY)
                    inline(p, txt, bold=True, colour="FFFFFF")
                else:
                    if i % 2 == 0:
                        shade(c, "F4F6F8")
                    inline(p, txt)
        # repeat header row
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
        self.after_table(container)

    def figure(self, container, path, caption):
        src = (REWORK / path).resolve() if not path.startswith("../") else (REWORK / path).resolve()
        if src.suffix == ".svg":
            src = FIG / "png" / (src.stem + ".png")
        from PIL import Image
        w_px, h_px = Image.open(src).size
        width = min(TEXT_W, Cm(14))
        if h_px / w_px > 0.9:  # tall figures narrower
            width = Cm(10)
        p = container.add_paragraph(style="Figure")
        p.add_run().add_picture(str(src), width=width)
        if caption:
            cp = container.add_paragraph(style="Caption Text")
            m = re.match(r"(Figure \d+\.\d+)\s*[—-]\s*(.+)", caption)
            if m:
                _run(cp, m.group(1) + "  ", True, False, NAVY, None)
                inline(cp, m.group(2))
            else:
                inline(cp, caption)


# ---------- document assembly ----------

def new_section(doc, start_type=WD_SECTION.NEW_PAGE):
    s = doc.add_section(start_type)
    for pg in s._sectPr.findall(qn("w:pgNumType")):  # copied from previous section; would restart numbering
        s._sectPr.remove(pg)
    for c in s._sectPr.findall(qn("w:cols")):
        s._sectPr.remove(c)
    return s


def running_head(section, left, right):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    section.different_first_page_header_footer = True
    h = section.header.paragraphs[0]
    fmt_head(h, left, right)
    for foot in (section.footer, section.first_page_footer):
        f = foot.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = add_field(f, "PAGE", "1")
        r.font.name = SANS
        r.font.size = Pt(8.5)
        r.font.color.rgb = rgb(MUTED)
    section.first_page_header.paragraphs[0].text = ""


def fmt_head(p, left, right):
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, alignment=2)  # right
    for txt, bold in ((left, False), ("\t", False), (right, True)):
        r = p.add_run(txt)
        r.font.name = SANS
        r.font.size = Pt(8)
        r.bold = bold or None
        r.font.color.rgb = rgb(MUTED)
    para_border(p, "bottom", 4, "C9D1D9", 3)


def no_head(section):
    for hf in (section.header, section.footer, section.first_page_header, section.first_page_footer):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs:
            p.text = ""


def page_setup(section):
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(3.5)
    section.top_margin, section.bottom_margin = Cm(2.6), Cm(2.4)
    section.header_distance, section.footer_distance = Cm(1.3), Cm(1.2)


def heading1(doc, container, number, title):
    p = container.add_paragraph(style="Heading 1")
    if number:
        r = p.add_run(f"{number}")
        r.font.color.rgb = rgb(GOLD)
        p.add_run("   ")
    p.add_run(title)
    para_border(p, "bottom", 18, "F0B429", 10)
    return p


def render_markdown(doc, R, text, kind):
    """kind: 'front', 'chapter', 'glossary'"""
    lines = text.splitlines()
    i = 0
    container = doc
    section_name = ""
    para_buf = []

    def flush():
        nonlocal para_buf
        if para_buf:
            style = {"Answers and Rationales": "Answer"}.get(section_name)
            if container is not doc:
                style = "Box Text"
            p = container.add_paragraph(style=style) if style else container.add_paragraph()
            inline(p, " ".join(para_buf))
            para_buf = []

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
            i += 1
            continue
        if s.startswith("# "):
            flush()
            i += 1
            continue  # H1 handled by caller
        if s.startswith("## "):
            flush()
            if container is not doc:
                R.after_table(doc)
            container = doc
            section_name = s[3:].strip()
            if section_name in BOXED_SECTIONS:
                fill, colour = BOXES[section_name]
                border = (6, "C9D1D9") if section_name == "Learning Objectives" else None
                t, container = R.box(doc, section_name, fill, colour, border)
            else:
                doc.add_paragraph(section_name, style="Heading 2 Unlisted" if section_name in UNLISTED else "Heading 2")
            i += 1
            continue
        if s.startswith("### "):
            flush()
            container.add_paragraph(s[4:], style="Heading 3")
            i += 1
            continue
        if s.startswith(">"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            head = block[0]
            if head.startswith("**Through Four Lenses**"):
                items = []
                for b in block[1:]:
                    m = re.match(r"-\s*\*\*(.+?):\*\*\s*(.+)", b)
                    if m:
                        items.append((m.group(1), m.group(2)))
                R.lenses(container, items)
            else:
                m = re.match(r"\*\*(.+?):\*\*\s*(.*)", head)
                label, body = (m.group(1), m.group(2)) if m else ("Note", head)
                body = " ".join([body] + block[1:])
                R.callout(container, label, body)
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            R.md_table(container, rows)
            continue
        m = re.match(r"!\[(.*?)\]\((.+?)\)", s)
        if m:
            flush()
            if "cover" in m.group(2):
                i += 1
                continue
            cap = ""
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and re.match(r"\*Figure", lines[j].strip()):
                cap = lines[j].strip().strip("*")
                i = j
            R.figure(container, m.group(2), cap or m.group(1))
            i += 1
            continue
        m = re.match(r"\*\*(Q\d+)\.\*\*\s*(.+)", s)
        if m and section_name == "Self-Assessment":
            flush()
            p = container.add_paragraph(style="Question")
            _run(p, m.group(1) + "  ", True, False, NAVY, None)
            stem = m.group(2)
            lo = re.search(r"\s*(\[LO\d+(?:,\s*LO\d+)*\])\s*$", stem)
            if lo:
                stem = stem[: lo.start()]
            inline(p, stem)
            if lo:
                _run(p, "  " + lo.group(1)[1:-1], False, False, MUTED, 8)
            i += 1
            opts = []
            while i < len(lines) and re.match(r"[A-E]\)\s", lines[i].strip()):
                opts.append(lines[i].strip())
                i += 1
            for k, o in enumerate(opts):
                op = container.add_paragraph(style="Option")
                _run(op, o[0] + "\t", True, False, NAVY, None)
                inline(op, o[2:].strip())
                op.paragraph_format.tab_stops.add_tab_stop(Cm(1.2))
                if k == len(opts) - 1:
                    op.paragraph_format.keep_with_next = False
                    op.paragraph_format.space_after = Pt(6)
            continue
        m = re.match(r"[-*]\s+(.+)", s)
        if m:
            flush()
            style = "Box Text" if container is not doc else "List Bullet"
            p = container.add_paragraph(style="List Bullet")
            if container is not doc:
                p.paragraph_format.space_after = Pt(2)
            inline(p, m.group(1))
            i += 1
            continue
        m = re.match(r"(\d+)\.\s+(.+)", s)
        if m:
            flush()
            num, body = m.group(1), m.group(2)
            if section_name == "References":
                p = container.add_paragraph(style="Reference")
                _run(p, num + ".\t", False, False, MUTED, None)
                reference_runs(p, body)
            elif section_name == "Learning Objectives":
                lo = re.match(r"\[(LO\d+)\]\s*(.+)", body)
                p = container.add_paragraph(style="Box Text")
                p.paragraph_format.left_indent = Cm(1.1)
                p.paragraph_format.first_line_indent = Cm(-1.1)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.1))
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _run(p, (lo.group(1) if lo else num + ".") + "\t", True, False, GOLD, 8.5)
                inline(p, lo.group(2) if lo else body)
            else:
                p = container.add_paragraph(style="Box Text" if container is not doc else None)
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.first_line_indent = Cm(-0.7)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
                _run(p, num + ".\t", True, False, NAVY, None)
                inline(p, body)
            i += 1
            continue
        if kind == "glossary":
            p = container.add_paragraph(style="Glossary Entry")
            inline(p, s)
            i += 1
            continue
        if section_name == "Answers and Rationales" and s.startswith("**"):
            flush()
            p = container.add_paragraph(style="Answer")
            m = re.match(r"\*\*(.+?)\*\*\s*(.*)", s)
            _run(p, m.group(1) + " ", True, False, NAVY, None)
            inline(p, m.group(2))
            i += 1
            continue
        para_buf.append(s)
        i += 1
    flush()
    if container is not doc:
        R.after_table(doc)


def title_pages(doc):
    # cover: zero-margin section with full-bleed image
    s = doc.sections[0]
    page_setup(s)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(0)
    s.header_distance = s.footer_distance = Cm(0)
    no_head(s)
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(FIG / "png" / "cover.png"), width=Cm(21), height=Cm(29.69))
    # move the section break into the cover paragraph so no blank page follows
    s2 = new_section(doc)
    brk = doc.paragraphs[-1]._p
    sect = brk.pPr.find(qn("w:sectPr")) if brk.pPr is not None else None
    if sect is not None:
        p._p.get_or_add_pPr().append(sect)
        brk.getparent().remove(brk)
    s2 = doc.sections[-1]
    page_setup(s2)
    no_head(s2)
    page_numbering(s2, "lowerRoman", 1)

    # title page
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(150)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(BOOK)
    r.font.name, r.font.size, r.bold = SERIF_HEAD, Pt(30), True
    r.font.color.rgb = rgb(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(SUB)
    r.font.name, r.font.size, r.italic = SERIF_HEAD, Pt(16), True
    r.font.color.rgb = rgb(MUTED)
    para_border(p, "bottom", 18, "F0B429", 14)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(AUTHOR)
    r.font.name, r.font.size, r.bold = SERIF, Pt(13), True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("For students of medicine, pharmacy, physical therapy and the health sciences")
    r.font.name, r.font.size = SANS, Pt(9.5)
    r.font.color.rgb = rgb(MUTED)

    # edition / notice page
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    notice = [
        ("The patients in this book (Mrs. Amal Hassan, Mr. Karim Adel and Ms. Lina Osei) are fictional. "
         "Any resemblance to real people is coincidental.", False),
        ("This book is for education. It is not clinical guidance. Laws, device approvals and product performance "
         "change quickly; check current sources before relying on any specific figure.", False),
    ]
    p.paragraph_format.space_before = Pt(430)
    for k, (t, b) in enumerate(notice):
        q = doc.add_paragraph() if k else p
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = q.add_run(t)
        r.font.name, r.font.size, r.bold = SANS, Pt(8.5), b or None
        r.font.color.rgb = rgb(MUTED)
        q.paragraph_format.space_after = Pt(5)

    # contents
    p = doc.add_paragraph("Contents", style="Title Heading")
    p.paragraph_format.page_break_before = True
    para_border(p, "bottom", 18, "F0B429", 10)
    tp = doc.add_paragraph()
    add_field(tp, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the contents.")


def build():
    doc = Document()
    setup_styles(doc)
    R = Renderer(doc)
    title_pages(doc)

    # front matter (How to Use This Book …)
    front = (REWORK / "00-front-matter.md").read_text(encoding="utf-8")
    front = front.split("\n## How to Use This Book", 1)[1]
    s = new_section(doc)
    page_setup(s)
    page_numbering(s, "lowerRoman")
    running_head(s, BOOK, "How to Use This Book")
    p = doc.add_paragraph("How to Use This Book", style="Heading 1")
    para_border(p, "bottom", 18, "F0B429", 10)
    render_markdown(doc, R, front, "chapter")

    chapters = sorted(REWORK.glob("ch*.md"))
    part_titles = {}
    for f in chapters:
        n = int(re.match(r"ch(\d+)", f.name).group(1))
        t = re.search(r"^# Chapter \d+:\s*(.+)$", f.read_text(encoding="utf-8"), re.M).group(1)
        part_titles[n] = t

    first = True
    for f in chapters:
        text = f.read_text(encoding="utf-8")
        n = int(re.match(r"ch(\d+)", f.name).group(1))
        title = re.search(r"^# Chapter \d+:\s*(.+)$", text, re.M).group(1)
        if n in PARTS:
            label, ptitle, span = PARTS[n]
            s = new_section(doc)
            page_setup(s)
            no_head(s)
            if first:
                page_numbering(s, "decimal", 1)
                first = False
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(200)
            r = p.add_run(label)
            r.font.color.rgb = rgb(GOLD)
            p.add_run("   " + ptitle)
            para_border(p, "bottom", 18, "F0B429", 10)
            q = doc.add_paragraph()
            r = q.add_run(span)
            r.font.name, r.font.size = SANS, Pt(9.5)
            r.font.color.rgb = rgb(MUTED)
            q.paragraph_format.space_after = Pt(14)
            k = n
            while k in part_titles and (k == n or k not in PARTS):
                q = doc.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.LEFT
                q.paragraph_format.tab_stops.add_tab_stop(Cm(1))
                r = q.add_run(f"{k}\t")
                r.bold, r.font.color.rgb = True, rgb(GOLD)
                r = q.add_run(part_titles[k])
                r.font.name, r.font.size = SERIF_HEAD, Pt(13)
                k += 1
        s = new_section(doc, WD_SECTION.NEW_PAGE)
        page_setup(s)
        running_head(s, BOOK, f"Chapter {n} · {title}")
        heading1(doc, doc, n, title)
        render_markdown(doc, R, text, "chapter")

    # glossary: heading single column, entries in two columns
    s = new_section(doc)
    page_setup(s)
    running_head(s, BOOK, "Glossary")
    heading1(doc, doc, None, "Glossary")
    s = new_section(doc, WD_SECTION.CONTINUOUS)
    page_setup(s)
    columns(s, 2)
    gl = (REWORK / "glossary.md").read_text(encoding="utf-8")
    render_markdown(doc, R, gl, "glossary")

    # hyphenation
    settings = doc.settings.element
    ah = OxmlElement("w:autoHyphenation")
    settings.append(ah)
    ul = OxmlElement("w:updateFields")
    ul.set(qn("w:val"), "true")
    settings.append(ul)

    core = doc.core_properties
    core.title = f"{BOOK}: {SUB}"
    core.author = AUTHOR
    core.language = "en-GB"
    core.comments = ""  # python-docx default is "generated by python-docx"
    core.last_modified_by = AUTHOR
    doc.save(OUT)
    print("wrote", OUT.name)


def word_finish(docx_path):
    """Open in Word: build contents, style it, embed fonts, save, export PDF."""
    import win32com.client as win32
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(docx_path))
        for lvl, (font, size, bold, before) in {1: (SANS, 10.5, True, 8), 2: (SERIF, 10, False, 0)}.items():
            st = d.Styles(-19 - lvl)  # wdStyleTOC1 = -20, TOC2 = -21
            st.Font.Name = font
            st.Font.Size = size
            st.Font.Bold = bold
            st.Font.Color = int(NAVY[4:6] + NAVY[2:4] + NAVY[0:2], 16) if lvl == 1 else int(INK[4:6] + INK[2:4] + INK[0:2], 16)
            st.ParagraphFormat.SpaceBefore = before
            st.ParagraphFormat.SpaceAfter = 2
            st.ParagraphFormat.LeftIndent = 0 if lvl == 1 else 28
        d.TablesOfContents(1).Update()
        d.Fields.Update()
        d.TablesOfContents(1).Update()
        d.EmbedTrueTypeFonts = True
        d.SaveSubsetFonts = False
        d.Save()
        pdf = docx_path.with_suffix(".pdf")
        try:
            open(pdf, "ab").close()
        except PermissionError:  # PDF open in a viewer
            pdf = pdf.with_name(pdf.stem + " (new).pdf")
            print("PDF is open elsewhere; writing", pdf.name)
        pdf = str(pdf)
        # 17 = wdExportFormatPDF, 0 = print quality, 1 = heading bookmarks
        d.ExportAsFixedFormat(pdf, 17, False, 0, 0, 1, 1, 0, True, True, 1, True, True, False)
        pages = d.ComputeStatistics(2)
        d.Close(False)
        print("wrote", pathlib.Path(pdf).name, f"({pages} pages)")
    finally:
        word.Quit()


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    build()
    if "--no-pdf" not in sys.argv:
        word_finish(OUT)
