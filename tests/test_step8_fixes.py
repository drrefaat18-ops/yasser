# tests/test_step8_fixes.py
"""Regression tests for the STEP 8 Codex review findings (docs/harness/reviews/step8-fixes.md)."""
import json, pathlib, shutil, subprocess, sys, tempfile, unittest
from harness.stages import complete_checks as cc
from harness.tools import build_book, check_book, config, convert_docx
from tests.helpers import REPO, stamp_project, temp_repo

NO_ERRATA = REPO / "tests/fixtures/positive-config/no-errata"


def copy_project(src=NO_ERRATA):
    tmp = pathlib.Path(tempfile.mkdtemp())
    dst = tmp / src.name
    shutil.copytree(src, dst)
    return tmp, dst


def edit(path, fn):
    d = json.loads(path.read_text(encoding="utf-8"))
    fn(d)
    path.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")


class PathContainment(unittest.TestCase):                                     # S8-01
    CASES = [("template.json", lambda t: t["paths"].update(chapters="../elsewhere"), "template.paths.chapters"),
             ("template.json", lambda t: t["paths"].update(normalized_source="C:/x.md"), "template.paths.normalized_source"),
             ("template.json", lambda t: t["paths"].update(allowed_asset_roots=["chapters", "../.."]), "allowed_asset_roots"),
             ("template.json", lambda t: t["paths"].update(glossary="../../g.md"), "template.paths.glossary"),
             ("design/chapter-plan.json", lambda c: c["chapters"][0].update(file="../../ch.md"), "chapter-plan ch01"),
             ("theme.json", lambda t: t["output"].update(basename="../out"), "theme.output.basename"),
             ("theme.json", lambda t: t["cover"].update(asset="/etc/cover.png"), "theme.cover.asset"),
             ("brief.json", lambda b: b["source"]["files"][0].update(path="../../book.docx"), "brief.source.files")]

    def test_every_configured_path_must_stay_inside(self):
        for rel, fn, needle in self.CASES:
            with self.subTest(needle):
                tmp, p = copy_project()
                try:
                    edit(p / rel, fn)
                    with self.assertRaises(config.ConfigError) as e:
                        config.load(p)
                    self.assertIn(needle, str(e.exception))
                    self.assertTrue(any(needle in x for x in config.problems(p)))
                finally:
                    shutil.rmtree(tmp)

    def test_intake_refuses_it(self):
        tmp, p = copy_project(REPO / "projects/ai-in-medicine")
        try:
            edit(p / "template.json", lambda t: t["paths"].update(chapters="../../escape"))
            _, probs = cc.intake(p)
            self.assertTrue(any(x.startswith("rule 9: template.paths.chapters") for x in probs), probs)
        finally:
            shutil.rmtree(tmp)

    def test_gated_tool_refuses_before_writing(self):
        with temp_repo(f"positive-config/{NO_ERRATA.name}", slug=NO_ERRATA.name, stamp=True) as root:
            p = root / "projects" / NO_ERRATA.name
            outside = root / "outside.md"
            shutil.copyfile(p / "chapters/ch01-forces.md", outside)
            edit(p / "design/chapter-plan.json", lambda c: c["chapters"][0].update(file="../../../outside.md"))
            stamp_project(root, p)   # even an approved plan cannot point outside
            before = outside.read_bytes()
            r = subprocess.run([sys.executable, "harness/tools/renumber_refs.py", "--project", f"projects/{p.name}", "--chapter", "ch01"],
                               cwd=root, capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 2, r.stderr)
            self.assertIn("ERROR CONFIG", r.stderr)
            self.assertEqual(outside.read_bytes(), before)


class FeatureFields(unittest.TestCase):                                       # S8-05
    CASES = [lambda t: t["assessment"]["mcq"].update(max_run=None), lambda t: t["assessment"]["mcq"].update(key_balance=None),
             lambda t: t["assessment"]["mcq"].update(count=None), lambda t: t["learning_objectives"].update(min=None),
             lambda t: t["assessment"]["case_question"].update(label=None), lambda t: t["perspectives"].update(balance_tolerance=None),
             lambda t: t["glossary"].update(minimum_terms=None), lambda t: t["errata"].update(status_column=None)]

    def test_enabled_feature_without_its_field_is_a_config_error(self):
        for n, fn in enumerate(self.CASES):
            with self.subTest(n):
                t = json.loads((NO_ERRATA / "template.json").read_text(encoding="utf-8"))
                if n == 7:
                    t["errata"] = {"enabled": True, "status_column": "Status", "open_values": ["open"], "closed_values": []}
                    t["paths"]["errata"] = "errata.md"
                fn(t)
                self.assertTrue(config.feature_problems(t))

    def test_cli_exit_2_not_a_traceback(self):
        with temp_repo(f"positive-config/{NO_ERRATA.name}", slug=NO_ERRATA.name, stamp=True) as root:
            p = root / "projects" / NO_ERRATA.name
            edit(p / "template.json", lambda t: t["assessment"]["mcq"].update(max_run=None))
            stamp_project(root, p)
            r = subprocess.run([sys.executable, "harness/tools/check_book.py", "--project", f"projects/{p.name}"],
                               cwd=root, capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 2)
            self.assertIn("ERROR CONFIG: assessment.mcq.max_run", r.stderr)
            self.assertNotIn("Traceback", r.stderr)


class CalloutMarker(unittest.TestCase):                                       # S8-03
    def test_mention_in_prose_is_not_a_callout(self):
        cfg = config.load(NO_ERRATA)
        ch = (NO_ERRATA / "chapters/ch01-forces.md").read_text(encoding="utf-8")
        mutated = "\n".join(l for l in ch.split("\n") if not l.startswith("> **Key Idea:**"))
        mutated = mutated.replace("Forces change how things move.", "Forces change how things move (see > **Key Idea:** boxes).")
        with tempfile.TemporaryDirectory() as t:
            f = pathlib.Path(t, "ch01.md")
            f.write_text(mutated, encoding="utf-8")
            rep = {c["id"]: c for c in check_book.check_chapter(f, cfg, {"id": "ch01", "word_budget": 0})["checks"]}
        self.assertEqual(rep["TPL-CALLOUT-MISSING"]["status"], "fail")

    def test_case_question_marker_must_start_a_line(self):
        cfg = config.load(NO_ERRATA)
        ch = (NO_ERRATA / "chapters/ch01-forces.md").read_text(encoding="utf-8")
        mutated = ch.replace("**Case Question.** Sami", "Think of a **Case Question.** Sami")
        with tempfile.TemporaryDirectory() as t:
            f = pathlib.Path(t, "ch01.md")
            f.write_text(mutated, encoding="utf-8")
            rep = {c["id"]: c for c in check_book.check_chapter(f, cfg, {"id": "ch01", "word_budget": 0})["checks"]}
        self.assertEqual(rep["MCQ-CASE"]["status"], "fail")


class FrontBudget(unittest.TestCase):                                         # S8-06
    def test_missing_front_matter_never_passes_its_budget(self):
        tmp, p = copy_project()
        try:
            (p / "chapters/00-front.md").unlink()
            rep, _ = check_book.check_book_level(config.load(p), [100])
            c = {x["id"]: x for x in rep["checks"]}
            self.assertEqual(c["BOOK-FRONT-MISSING"]["status"], "fail")
            self.assertEqual(c["BUDGET-FRONT"]["status"], "not_applicable")
            self.assertIn("missing", c["BUDGET-FRONT"]["message"])
        finally:
            shutil.rmtree(tmp)


class FidelityText(unittest.TestCase):                                        # S8-02
    def test_annotation_text_does_not_hide_lost_body_text(self):
        from docx import Document
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        spec = {"source": pathlib.Path(REPO / "tests/fixtures/ingest/make_fixtures.py")}
        import importlib.util
        m = importlib.util.spec_from_file_location("make_fixtures", spec["source"])
        mk = importlib.util.module_from_spec(m)
        m.loader.exec_module(mk)
        long = "Mass and force and motion. " * 12
        d = Document()
        d.add_heading("Chapter 1: Motion", level=1)
        d.add_paragraph(mk.TEXT)
        p = d.add_paragraph("A boxed note follows.")
        p._p.append(parse_xml(f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox>'
                              f'<w:txbxContent><w:p><w:r><w:t>{long}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
        q = d.add_paragraph("A sourced claim.")
        q.add_run()._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="1"/>'))
        with tempfile.TemporaryDirectory() as t:
            f = pathlib.Path(t, "x.docx")
            d.save(f)
            mk.FOOTNOTES_XML = mk.FOOTNOTES_XML.replace("Newton 1687.", long)
            mk.add_footnotes_part(f)
            cfg = config.load(NO_ERRATA)
            cfg["template"]["ingest"] = {"heading_rules": [], "toc_start_pattern": None, "toc_entry_pattern": None,
                                         "toc_end_pattern": None, "drop_source_toc": False}
            _, _, rep = convert_docx.convert(f, cfg)
        self.assertEqual(rep["structures"]["footnotes"]["status"], "lossy")
        self.assertEqual(rep["structures"]["text"]["status"], "lost")

    def test_output_text_excludes_inlined_annotations(self):
        self.assertEqual(convert_docx.output_counts("abc [footnote: long text here]\n[equation: x=1]\n")["text"], 3)


class BuilderLabels(unittest.TestCase):                                       # S8-04
    def test_display_labels_come_from_theme_and_template(self):
        tmp, p = copy_project()
        try:
            edit(p / "theme.json", lambda t: t["labels"].update(question_prefix="Question ", objective_prefix="Goal "))
            edit(p / "template.json", lambda t: t["assessment"]["mcq"].update(option_display_labels=["a", "b", "c", "d"]))
            from docx import Document
            doc = Document(str(build_book.build(config.load(p))))
            texts = [x.text for x in doc.paragraphs] + [x.text for t in doc.tables for row in t.rows for c in row.cells
                                                         for x in c.paragraphs]
            self.assertTrue(any(t.startswith("Question 1  ") for t in texts), texts)
            self.assertTrue(any(t.startswith("Goal 1\t") for t in texts))
            self.assertTrue(any(t.startswith("a\t") for t in texts))
            self.assertTrue(any(t.startswith("Question 2. a ") for t in texts))
            self.assertFalse(any(t.startswith(("Q1 ", "LO1\t", "A\t")) for t in texts))
        finally:
            shutil.rmtree(tmp)


if __name__ == "__main__":
    unittest.main()
