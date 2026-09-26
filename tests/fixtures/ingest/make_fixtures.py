"""Generate the DOCX ingest fixtures (plan Task 8.4), one high-risk structure each. Run once; the .docx files are
committed with hand-written *.expected.json beside them.

Usage: python tests/fixtures/ingest/make_fixtures.py
"""
import io, pathlib, re, struct, zipfile, zlib
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm

HERE = pathlib.Path(__file__).resolve().parent
TEXT = "Motion changes when a force acts. A larger mass needs a larger force for the same change."


def png(r, g, b):
    raw = b"".join(b"\x00" + bytes([r, g, b]) * 4 for _ in range(3))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 3, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def base():
    d = Document()
    d.core_properties.author = "fixture"
    d.add_heading("Chapter 1: Motion", level=1)
    d.add_paragraph(TEXT)
    return d


def headings():
    d = base()
    d.add_heading("Forces", level=2)
    d.add_paragraph(TEXT)
    d.add_heading("Friction", level=3)
    d.add_paragraph(TEXT)
    d.add_paragraph().add_run("Summary Notes").bold = True     # matched by template.ingest.heading_rules
    d.add_paragraph(TEXT)
    return d


def tables():
    d = base()
    for rows, cols in ((2, 2), (3, 3)):
        t = d.add_table(rows=rows, cols=cols)
        for i in range(rows):
            for j in range(cols):
                t.cell(i, j).text = f"r{i}c{j}"
        d.add_paragraph(TEXT)
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Key idea: force changes motion."
    d.add_paragraph(TEXT)
    return d


def figures():
    d = base()
    for rgb in ((255, 0, 0), (0, 0, 255)):
        d.add_picture(io.BytesIO(png(*rgb)), width=Cm(2))
        d.add_paragraph(TEXT)
    return d


def equations():
    d = base()
    p = d.add_paragraph("The energy of a mass is ")
    p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>E=mc2</m:t></m:r></m:oMath>'))
    d.add_paragraph(TEXT)
    return d


def footnotes():
    d = base()
    for fid in (1, 2):
        p = d.add_paragraph(f"A claim that needs a source number {fid}.")
        r = p.add_run()
        r._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="{fid}"/>'))
    d.add_paragraph(TEXT)
    return d


FOOTNOTES_XML = ('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:id="1"><w:p><w:r><w:t>Newton 1687.</w:t></w:r></w:p></w:footnote>'
                 '<w:footnote w:id="2"><w:p><w:r><w:t>Halliday 2013.</w:t></w:r></w:p></w:footnote></w:footnotes>')


def add_footnotes_part(path):
    """python-docx has no footnote API: add word/footnotes.xml, its relationship and content type."""
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    items["word/footnotes.xml"] = FOOTNOTES_XML.encode()
    rels = items["word/_rels/document.xml.rels"].decode()
    items["word/_rels/document.xml.rels"] = rels.replace("</Relationships>", '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
                                                         'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>').encode()
    ct = items["[Content_Types].xml"].decode()
    items["[Content_Types].xml"] = ct.replace("</Types>", '<Override PartName="/word/footnotes.xml" ContentType="application/'
                                              'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>').encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in items.items():
            z.writestr(n, data)


def toc():
    d = Document()
    d.add_paragraph("Contents")
    d.add_paragraph("Chapter 1: Motion\t3")
    d.add_paragraph("Chapter 2: Energy\t9")
    d.add_heading("Chapter 1: Motion", level=1)
    d.add_paragraph(TEXT)
    d.add_heading("Chapter 2: Energy", level=1)
    d.add_paragraph(TEXT)
    return d


def dropped_table():
    d = base()
    p = d.add_paragraph("A diagram sits in a text box here.")
    cell = lambda t: f"<w:tc><w:p><w:r><w:t>{t}</w:t></w:r></w:p></w:tc>"
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="width:200pt;height:60pt">'
        f'<v:textbox><w:txbxContent><w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>'
        f'<w:tr>{cell("mass")}{cell("force")}</w:tr></w:tbl><w:p/></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
    d.add_paragraph(TEXT)
    return d


FIXTURES = {"headings": headings, "tables": tables, "figures": figures, "equations": equations,
            "footnotes": footnotes, "toc": toc, "dropped-table": dropped_table}


if __name__ == "__main__":
    for name, fn in FIXTURES.items():
        out = HERE / f"{name}.docx"
        fn().save(out)
        if name == "footnotes":
            add_footnotes_part(out)
        print("wrote", out.name)
