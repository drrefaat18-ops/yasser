"""build_slides: one deck per chapter, driven only by project config (Rule 8)."""
import pathlib, shutil, tempfile, unittest
from pptx import Presentation
from harness.tools import build_slides
from tests.helpers import REPO

FIXTURE = REPO / "tests" / "fixtures" / "editorial-book"


class BuildSlidesTest(unittest.TestCase):
    def test_one_deck_per_chapter_with_questions_and_notes(self):
        with tempfile.TemporaryDirectory() as tmp:
            proj = pathlib.Path(tmp) / "book"
            shutil.copytree(FIXTURE, proj)                 # fixtures are immutable; work on a copy
            self.assertEqual(build_slides.main(proj), 0)
            decks = sorted((proj / "slides").glob("*.pptx"))
            self.assertEqual([d.stem for d in decks], ["ch01-forces", "ch02-friction-graphs"])
            prs = Presentation(decks[0])
            texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
            self.assertIn("Forces and Motion", texts[0])                     # title slide
            self.assertTrue(any("Your turn" in t for t in texts))            # question slide
            self.assertTrue(any("Check your answer" in t for t in texts))    # answer slide
            self.assertTrue(any(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in prs.slides))
            self.assertTrue((proj / "slides" / "_assets" / "logo.png").is_file())   # logo badge from theme

    def test_fit_to_minutes_timings_add_up(self):
        with tempfile.TemporaryDirectory() as tmp:
            proj = pathlib.Path(tmp) / "book"
            shutil.copytree(FIXTURE, proj)
            (proj / "slides.json").write_text('{"minutes": 12, "fit_to_minutes": true}', encoding="utf-8")
            self.assertEqual(build_slides.main(proj), 0)
            prs = Presentation(proj / "slides" / "ch01-forces.pptx")
            last = prs.slides[-1].notes_slide.notes_text_frame.text
            self.assertIn("at 12:00 of 12:00", last)                         # notes add up to the session
            titles = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
            self.assertFalse(any(t.endswith("(1/2)") for t in titles if t.replace("(1/2)", "(2/2)") not in titles))

    def test_glass_look_same_slides_with_background(self):
        with tempfile.TemporaryDirectory() as tmp:
            proj = pathlib.Path(tmp) / "book"
            shutil.copytree(FIXTURE, proj)
            self.assertEqual(build_slides.main(proj), 0)
            flat = len(Presentation(proj / "slides" / "ch01-forces.pptx").slides)
            (proj / "slides.json").write_text('{"look": "glass"}', encoding="utf-8")
            self.assertEqual(build_slides.main(proj), 0)
            prs = Presentation(proj / "slides" / "ch01-forces.pptx")
            self.assertEqual(len(prs.slides), flat)                          # content unchanged
            self.assertTrue(all(s.shapes[0].shape_type == 13 for s in prs.slides))   # blurred field behind each slide
            whys = [sh.text_frame.text for s in prs.slides for sh in s.shapes
                    if sh.has_text_frame and sh.text_frame.text.startswith("Why:")]
            self.assertTrue(whys and all(len(w) > 8 for w in whys))         # the reason follows "Why:"

    def test_mix_and_parse_helpers(self):
        self.assertEqual(build_slides.mix("000000", "FFFFFF", 1), "FFFFFF")
        self.assertEqual(build_slides.clean("Risk fell [2, 3] to 1^+^"), "Risk fell to 1+")


if __name__ == "__main__":
    unittest.main()
