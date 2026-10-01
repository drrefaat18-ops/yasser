"""build_study: source-faithful study decks; the fidelity check refuses anything not in the chapter."""
import json, pathlib, re, shutil, tempfile, unittest
from pptx import Presentation
from harness.tools import build_study
from tests.helpers import REPO

FIXTURE = REPO / "tests" / "fixtures" / "editorial-book"


def outline(md, extra=None):
    """One points slide per section: the first sentence of its first paragraph, quoted as its own source."""
    secs = []
    for sec, (title, prose, _) in build_study.sections_of(md).items():
        first = re.split(r"(?<=[.!?])\s", re.sub(r"\s*\[\d+\]", "", prose[0]))[0]
        secs.append({"sec": sec, "slides": [{"kind": "points", "title": title, "points": [{"text": first, "src": first}]}]})
    if extra:
        secs[0]["slides"][0]["points"].append(extra)
    return {"chapter": "ch01-forces.md", "sections": secs}


class BuildStudyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.proj = pathlib.Path(self.tmp.name) / "book"
        shutil.copytree(FIXTURE, self.proj)                 # fixtures are immutable; work on a copy
        self.md = (self.proj / "chapters" / "ch01-forces.md").read_text(encoding="utf-8")
        (self.proj / "slides" / "study").mkdir(parents=True)

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, data):
        (self.proj / "slides" / "study" / "ch01-forces.json").write_text(json.dumps(data), encoding="utf-8")

    def test_faithful_outline_builds_a_deck(self):
        self.write(outline(self.md))
        self.assertEqual(build_study.main(self.proj), 0)
        deck = self.proj / "slides" / "study" / "out" / "ch01-forces-study.pptx"
        cells = lambda sh: [c.text for r in sh.table.rows for c in r.cells] if sh.has_table else []
        texts = [" ".join([sh.text_frame.text for sh in s.shapes if sh.has_text_frame] + [t for sh in s.shapes for t in cells(sh)])
                 for s in Presentation(deck).slides]
        self.assertTrue(any("Key terms" in t and "Mass" in t for t in texts))       # glossary definitions
        self.assertTrue(any("Questions 1–2" in t for t in texts))                   # two questions per slide
        self.assertIn("Answer key", texts[-1])                                      # all answers at the end
        self.assertTrue(any("Key takeaways" in t for t in texts))
        self.assertFalse((self.proj / "slides" / "study" / "out" / "ch02-friction-graphs-study.pptx").exists())

    def test_words_from_outside_the_chapter_are_refused(self):
        self.write(outline(self.md, {"text": "Quantum entanglement drives every trolley.", "src": "Quantum entanglement"}))
        self.assertEqual(build_study.main(self.proj), 1)
        report = json.loads((self.proj / "slides" / "study" / "out" / "study-qa.json").read_text(encoding="utf-8"))
        problems = " ".join(x["problem"] for x in report["ch01-forces"]["fidelity"])
        self.assertIn("not verbatim", problems)
        self.assertIn("quantum", problems)
        self.assertFalse((self.proj / "slides" / "study" / "out" / "ch01-forces-study.pptx").exists())

    def test_diagram_nodes_are_checked_and_drawn(self):
        data = outline(self.md)
        first = data["sections"][0]["slides"][0]["points"][0]
        data["sections"][0]["slides"].append({"kind": "flow", "title": "Force and Motion", "nodes": [
            {"label": "Force", "text": first["text"], "src": first["src"]},
            {"label": "Motion", "text": first["text"], "src": first["src"]}],
            "caption": [{"text": first["text"], "src": first["src"]}]})
        self.write(data)
        self.assertEqual(build_study.main(self.proj), 0)
        data["sections"][0]["slides"][-1]["nodes"].append({"label": "Warp drive", "text": "Spaceships.", "src": "Spaceships"})
        self.write(data)
        self.assertEqual(build_study.main(self.proj), 1)                            # invented node refused

    def test_stats_icons_and_variety(self):
        data = outline(self.md)
        first = data["sections"][0]["slides"][0]["points"][0]
        node = {"label": first["text"].split()[0], "text": first["text"], "icon": "🦠", "src": first["src"]}
        data["sections"][0]["slides"] += [{"kind": "stats", "title": "Force and Motion", "nodes": [node, dict(node)]},
                                          {"kind": "stats", "title": "Force and Motion", "nodes": [node]}]
        self.write(data)
        self.assertEqual(build_study.main(self.proj), 0)
        self.assertTrue(any("stats layout" in v for v in build_study.variety(data)))   # stats twice in a row

    def test_missing_section_is_reported(self):
        data = outline(self.md); data["sections"] = data["sections"][:1]
        issues = build_study.check(data, self.md, build_study.sections_of(self.md))
        self.assertTrue(any(p == "no slide" for _, p in issues))

    def test_topics_split_a_source_at_chapter_headings(self):
        import re
        rx = re.compile(r"^# Topic (?P<num>\d+): (?P<title>.+)$")
        parts = build_study.topics("# Topic 1: A\n## 1.1 X\ntext\n# Topic 2: B\nmore", rx)
        self.assertEqual(sorted(parts), ["1", "2"])
        self.assertIn("## 1.1 X", parts["1"]); self.assertNotIn("more", parts["1"])

    def test_norm_ignores_markup_and_quote_style(self):
        self.assertEqual(build_study.norm("**Need**  is “felt”"), 'need is "felt"')
        self.assertEqual(build_study.stem("studies"), "study")


if __name__ == "__main__":
    unittest.main()
